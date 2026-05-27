from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Eratex_Planning_Scheduling_Tool_BRD_Professional_Aligned.docx"

NAVY = "183A5F"
BLUE = "2E75B6"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8F1F8"
LIGHT_GRAY = "F2F4F7"
GRID = "D8E2EC"
INK = "172033"
MUTED = "657080"
WHITE = "FFFFFF"
GREEN = "2E7D32"
AMBER = "B7791F"
RED = "B42318"


def rgb(hex_value):
    return RGBColor.from_string(hex_value.replace("#", ""))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_border(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color.replace("#", ""))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def paragraph_border(paragraph, color=BLUE, size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color.replace("#", ""))
    p_bdr.append(bottom)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_run(p, text, bold=False, italic=False, color=None, size=None):
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_para(doc, text="", style=None, before=None, after=None, keep=False, align=None, color=None, bold=False, size=None):
    p = doc.add_paragraph(style=style)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    if align is not None:
        p.alignment = align
    add_run(p, text, bold=bold, color=color, size=size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)


def add_callout(doc, title, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, "C9DCEB")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    add_run(p, f"{title}: ", bold=True, color=DARK_BLUE)
    add_run(p, text)
    return table


def add_table(doc, headers, rows, widths_dxa=None, font_size=8.8):
    widths_dxa = widths_dxa or [int(9360 / len(headers))] * len(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_run(p, str(text), bold=True, color=WHITE, size=font_size)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, "FFFFFF" if ridx % 2 == 0 else "EEF4F9")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_run(p, str(text), size=font_size)
    add_para(doc, "", after=4)
    return table


def add_key_value_table(doc, rows):
    return add_table(doc, ["Field", "Value"], rows, [2100, 7260], font_size=9)


def part(doc, title):
    p = add_para(doc, title, before=18, after=10, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.style = doc.styles["Heading 1"]
    for run in p.runs:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = rgb(NAVY)
    paragraph_border(p, BLUE, "6")


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Eratex BRD | Planning & Scheduling Tool | Aligned Revision"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(MUTED)
    paragraph_border(hp, GRID, "4")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "Confidential - Eratex Internal Use Only"
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in fp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(MUTED)
    pn = footer.add_paragraph()
    pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(pn)
    for run in pn.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(MUTED)


def cover(doc):
    add_para(doc, "BUSINESS REQUIREMENTS DOCUMENT", before=80, after=4, align=WD_ALIGN_PARAGRAPH.CENTER, color=WHITE, bold=True, size=22)
    last = doc.paragraphs[-1]
    last.paragraph_format.space_before = Pt(90)
    last.paragraph_format.space_after = Pt(0)
    last._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), NAVY)
    last._p.get_or_add_pPr().append(shading)

    p = add_para(doc, "Eratex End-to-End Planning & Scheduling Tool", after=2, align=WD_ALIGN_PARAGRAPH.CENTER, color=WHITE, size=18)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), NAVY)
    p._p.get_or_add_pPr().append(shading)

    p = add_para(doc, "Aligned with Rajesh Inputs and BRD Diff Recommendations", after=24, align=WD_ALIGN_PARAGRAPH.CENTER, color="C9DCEB", size=12)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), NAVY)
    p._p.get_or_add_pPr().append(shading)

    add_key_value_table(doc, [
        ("Company", "Eratex"),
        ("Document Type", "Business Requirements Document"),
        ("Product Scope", "Denim-aware planning and scheduling platform for garment manufacturing operations"),
        ("Manufacturing Scope", "Denim bottoms and chinos; denim treated as dominant complexity driver"),
        ("Baseline", "Eratex_Planning_Scheduling_Tool_BRD_Professional_Generated.docx"),
        ("Alignment Source", "04_Diff_Eratex_BRD_Refinement_From_Rajesh_Inputs.md"),
        ("Version", "1.1 aligned draft"),
        ("Prepared Date", "2026-05-27"),
        ("Status", "Draft for business review"),
    ])
    add_callout(doc, "Core revision", "The BRD now frames the product as a governed, denim-aware operating spine rather than only a scheduling screen or Excel replacement.")
    doc.add_page_break()


def build():
    doc = Document()
    configure_doc(doc)
    cover(doc)

    add_para(doc, "1. Executive Summary", style="Heading 1")
    add_para(doc, "Eratex operates as an export-oriented garment manufacturer serving international customers, primarily across denim bottoms and chinos. The production environment is not a simple cut-and-sew flow. For denim, commercial value is often created through dry processing, wet washing, fading, distressing, tearing, touch-up, repeat wash cycles, and buyer-specific shade/effect outcomes.")
    add_para(doc, "Although a planning tool such as FastReact may exist in the operating environment, planning still appears to depend heavily on Excel and manual coordination. The deeper issue is not only tool adoption. The deeper issue is that the real operating complexity of denim production - especially wash/laundry, rework, WIP movement, capacity changes, order changes, and shipment-risk recovery - is not governed by one trusted, live operating spine.")
    add_para(doc, "The proposed solution should therefore be an end-to-end, denim-aware planning and scheduling platform that converts customer demand into executable production control. It must support order visibility, PCD readiness, capacity-aware planning, daily release, sewing line loading, wash and rewash planning, WIP ageing, exceptions, recovery actions, and shipment readiness.")
    add_callout(doc, "Business outcome", "Maintain high OTIF, improve utilization, reduce overtime and firefighting, reduce Excel dependency, and protect shipment commitments through early, constraint-aware control.")

    add_para(doc, "2. Business Context", style="Heading 1")
    add_para(doc, "2.1 Export-Oriented Operating Context", style="Heading 2")
    add_para(doc, "Eratex should be treated as an export-oriented garmenting operation where customer delivery commitments, buyer approvals, inspection readiness, documentation, packing, and dispatch reliability are central to business performance. Orders are expected to be driven by international customers, with US and Europe referenced as primary markets in the consultant input.")
    add_para(doc, "2.2 End-to-End Garment Flow", style="Heading 2")
    add_bullets(doc, [
        "Customer enquiry, quotation visibility, sampling, buyer approval, and order confirmation.",
        "Fabric and trim procurement, fabric inward, fabric QC, shade-lot control, and PCD readiness.",
        "Cutting, sewing, dry process, wet wash, repeat wash or rework, finishing, final QC, packing, documentation, and shipment.",
    ])
    add_para(doc, "2.3 Denim-Specific Manufacturing Complexity", style="Heading 2")
    add_para(doc, "Denim manufacturing has a different complexity profile from basic garmenting. Visible effects such as fading, distressing, tearing, whiskering, scraping, grinding, uneven shade, hand feel, and fashion-wash outcomes may be intentional value-add features rather than defects.")
    add_para(doc, "This makes wash/laundry a value-creation, quality-risk, and capacity-constraining production domain. Dry process, wet wash, repeat wash cycles, post-wash shade/effect checks, and rewash loops must be planned explicitly because they affect capacity, WIP ageing, recovery cost, and shipment risk.")
    add_para(doc, "2.4 Chinos as Lower-Complexity Variant", style="Heading 2")
    add_para(doc, "Chinos should be supported through the same operating spine, but with simpler route and wash complexity unless business evidence proves otherwise.")

    add_para(doc, "3. Reconciled Problem Statement", style="Heading 1")
    add_para(doc, "Eratex is a large export-oriented denim and chinos garment manufacturer where production complexity extends beyond cutting and sewing into customer approvals, nominated material procurement, fabric QC, PCD readiness, sewing line loading, dry/wet wash processing, repeat wash cycles, quality holds, WIP movement, finishing, packing, and shipment readiness.")
    add_para(doc, "Although a formal planning tool may exist, the factory appears to rely heavily on Excel and manual coordination because the real operating complexity - especially denim wash/rework, shifting capacity constraints, order changes, and exception recovery - is not fully governed by one trusted system.")
    add_para(doc, "The factory is able to maintain OTIF above 95%, but this may be achieved through extra operating expense, overtime, expediting, resequencing, and manual firefighting, while resource utilization remains around 70%.")
    add_callout(doc, "Core problem", "The core problem is not the absence of software. It is the absence of a live, capacity-aware, wash-aware, WIP-trusted, exception-owned operating spine that converts customer demand into executable production control and reacts when factory reality changes.")
    add_para(doc, "3.1 Symptoms Observed in Current Planning", style="Heading 2")
    add_bullets(doc, [
        "Planning may be formally digital but operationally manual.",
        "OTIF is maintained through firefighting and cost absorption.",
        "Utilization is low because flow is unstable.",
        "Workcenter constraints shift without timely visibility.",
        "Wash rework and repeat cycles may not be fully capacity-planned.",
        "PCD readiness may not be enforced as a hard gate.",
        "WIP ageing may remain hidden between departments.",
        "Shipment readiness may be assessed too late.",
    ])
    add_para(doc, "3.2 Root Cause Behind the Symptoms", style="Heading 2")
    add_para(doc, "Excel remains active because the formal planning environment may not yet represent the real production states, wash/rework behaviour, capacity-change reactions, and exception ownership required by daily operations.")

    add_para(doc, "4. Business Need and Operating Spine", style="Heading 1")
    add_para(doc, "Eratex requires an integrated planning and scheduling tool that becomes the single operating layer for production control. It must answer daily readiness, bottleneck, WIP, exception, ownership, and shipment-risk questions.")
    add_table(doc, ["Business Question", "Required System Response"], [
        ("Which orders are ready or blocked?", "Show lifecycle, PCD, material, wash, quality, and shipment readiness status."),
        ("Where is the bottleneck?", "Show load, queue, WIP ageing, utilization, and constraint flags by workcenter."),
        ("What can be released today?", "Validate material, prior-process, capacity, quality, manpower, and next-process readiness."),
        ("What recovery action is required?", "Recommend rule-based actions with owner, due date, and approval requirement."),
        ("What changed in the plan?", "Maintain version, audit trail, boundary-case event, and impact preview."),
    ], [2600, 6760])
    add_para(doc, "The product must also define how the system reacts when operating reality changes:", style="Heading 2")
    add_bullets(doc, [
        "What happens if an order is cancelled?",
        "What happens if capacity is lost or added?",
        "What happens if wash must repeat?",
        "What happens if quality fails after wash?",
        "What happens if shipment date is pulled in?",
        "What happens if a frozen weekly plan needs to change?",
        "What is recommended, what is blocked, and what requires approval?",
    ])

    add_para(doc, "5. Business Objectives and Success Metrics", style="Heading 1")
    add_para(doc, "5.1 Primary Objectives", style="Heading 2")
    add_numbered(doc, [
        "Establish one integrated planning view from order confirmation to shipment.",
        "Enforce PCD readiness before cutting starts.",
        "Convert weekly production plans into daily executable releases.",
        "Monitor total workcenter load and identify shifting constraints.",
        "Improve sewing line loading using realistic capacity assumptions.",
        "Treat denim wash/laundry as a value-creation, quality-risk, and capacity-constraining production domain, while supporting chinos as a simpler configured route variant where applicable.",
        "Make WIP ageing visible across departments.",
        "Track exceptions with owner, severity, due date, and recovery action.",
        "Distinguish production completion from true shipment readiness.",
        "Reduce reliance on Excel-based planning and manual firefighting.",
        "Define scheduling boundary-case behaviour for cancellation, capacity change, rewash, quality failure, shipment pull-in, and plan change.",
        "Establish the MVP planning grain so business users, planners, and builders share the same level of scheduling detail.",
        "Produce a Scheduling Behaviour Rulebook before deep planning-engine implementation.",
    ])
    add_para(doc, "5.2 Success Metrics", style="Heading 2")
    add_table(doc, ["Metric", "Target Direction"], [
        ("OTIF", "Maintain above 95% while reducing recovery cost."),
        ("Utilization", "Improve from around 70% through better flow and readiness."),
        ("Planning adherence", "Improve weekly and daily plan compliance."),
        ("PCD readiness accuracy", "Reduce cutting releases with open blockers."),
        ("WIP ageing", "Reduce ageing at key handover points."),
        ("Wash queue ageing", "Reduce sewn garments waiting for wash."),
        ("Rewash visibility", "Capture 100% of rewash cycles with reason, quantity, and capacity impact."),
        ("Boundary-case response", "Cancellation, capacity change, and shipment pull-in events have visible owner and action."),
        ("Capacity definition completeness", "100% of MVP workcenters have agreed capacity unit, bucket, and recovery levers."),
        ("Rulebook completion", "Scheduling Behaviour Rulebook approved before full planning-engine build."),
    ], [2800, 6560])

    add_para(doc, "6. Product Scope and MVP Boundaries", style="Heading 1")
    add_para(doc, "The ten MVP surfaces remain valid. MVP scope should also capture upstream milestone visibility for customer enquiry, quotation status, sample submission, sample approval, technical file receipt, and order confirmation. Full quotation costing and commercial feasibility workflow should remain outside MVP unless confirmed by the business team.")
    add_para(doc, "6.1 Product Complexity Classes", style="Heading 2")
    add_table(doc, ["Complexity Class", "Planning Implication"], [
        ("DENIM_HEAVY_WASH", "Dry/wet route, fashion effects, higher rewash risk, shade/effect QC mandatory."),
        ("DENIM_BASIC_WASH", "Standard wash route, lower rewash probability."),
        ("CHINO_BASIC_WASH", "Simpler route, lower wash complexity."),
        ("CHINO_NO_COMPLEX_WASH", "Standard garmenting route without denim-style wash complexity."),
    ], [2700, 6660])
    add_para(doc, "6.2 MVP Surfaces", style="Heading 2")
    add_table(doc, ["#", "Surface", "Primary Role in MVP"], [
        ("1", "Order Lifecycle", "One digital thread and upstream milestone visibility."),
        ("2", "PCD Readiness", "Gate-based readiness before cutting."),
        ("3", "Weekly Planning Workbench", "Capacity-checked coordination contract."),
        ("4", "Daily Production Release", "Executable release based on true readiness."),
        ("5", "Workcenter Load Monitor", "Load, queue, capacity, WIP, and bottleneck visibility."),
        ("6", "Sewing Line Loading", "Realistic line targets using SMV, manpower, and net-good output."),
        ("7", "Wash Planning", "Wash/rework as value, quality, and capacity domain."),
        ("8", "WIP and Queue Monitoring", "Hidden waiting time and ownership made visible."),
        ("9", "Exception and Alert Surface", "Structured recovery action with owner and due date."),
        ("10", "Shipment Readiness", "Dispatch readiness separate from production completion."),
    ], [650, 3000, 5710])

    add_para(doc, "7. MVP Scheduling Grain Decision", style="Heading 1")
    add_callout(doc, "Decision", "MVP scheduling grain should be order x style x color/shade-lot x production stage x workcenter/line x date/shift, with wash handled at wash-batch grain.")
    add_para(doc, "Operation-level detail should exist in master data and line-balance review, but the first scheduling engine should not attempt full operation-by-operation finite scheduling unless validated as necessary.")

    add_para(doc, "8. Scheduling Boundary Cases and System Reaction Rules", style="Heading 1")
    add_para(doc, "The system must define what happens when orders are cancelled, quantities change, capacity is lost or added, wash repeats, quality fails, shipment dates move, or frozen plans need to change. MVP should provide rule-based impact preview and recommended recovery actions, with planner approval for committed schedule changes.")
    add_table(doc, ["Event", "Stage", "System Reaction", "Planner Action", "Audit / Approval"], [
        ("Order cancellation", "Before procurement", "Release planned capacity; cancel open plan intent.", "Review commercial impact.", "Audit."),
        ("Order cancellation", "After cutting", "Freeze or reclassify WIP; update load and shipment risk.", "Management disposition decision.", "Approval required."),
        ("Capacity loss", "Sewing", "Recalculate load, risk, and affected orders.", "Reassign, split, or recover.", "Audit."),
        ("Capacity loss", "Wash", "Recalculate wash queue, rewash reserve, and shipment risk.", "Resequence, overtime, split batch.", "Approval if committed."),
        ("Wash rework", "Post-wash QC", "Create rewash WIP and load.", "Approve and resequence.", "Audit."),
        ("Shipment pull-in", "Any stage", "Recalculate priority, capacity conflict, and recovery options.", "Prioritize, simulate, approve.", "Approval if override."),
        ("Frozen plan change", "Firm/frozen zone", "Create plan-change request and impact preview.", "Approve, reject, or revise.", "Approval required."),
    ], [1600, 1350, 2700, 2200, 1510], font_size=8.2)

    part(doc, "Part A: MVP Surface Requirements")
    surfaces = [
        ("9. Order Lifecycle Surface", "Provide one digital thread for every customer order from enquiry/quotation visibility through shipment.", [
            "Order master view with customer, buyer, style, PO, quantity, delivery date, current stage, risk, owner, and next action.",
            "Upstream milestone visibility: enquiry, quotation, sample request, sample submitted, sample approved, order confirmed.",
            "Order change status: cancellation, quantity change, delivery change, and split shipment request.",
            "Product complexity class, wash complexity class, and boundary-case flag.",
        ], [
            "Every order has one system-level lifecycle status.",
            "Delayed or changed orders have visible reason, owner, and next action.",
        ]),
        ("10. PCD Readiness Surface", "Enforce production readiness before cutting starts.", [
            "Validate PO, BOM, fabric receipt, fabric QC, shade lots, shrinkage report, trims, pattern, marker, PP sample, wash standard, line allocation, wash capacity, and QC file.",
            "Add readiness checks for product complexity class, wash route, dry/wet requirement, fashion-effect standard, post-wash QC criteria, rewash handling rule, and capacity impact.",
            "Support Ready, Conditionally Ready, Blocked, and Escalated statuses.",
        ], [
            "No order is released to cutting unless wash route and complexity class are visible.",
            "Conditional release includes reason, owner, expiry, and recovery action.",
        ]),
        ("11. Weekly Planning Workbench", "Convert the order book into an executable weekly cross-functional plan.", [
            "2-4 week firm plan and 8-12 week rolling visibility.",
            "Order sequencing, cutting plan, sewing allocation, wash plan, finishing, inspection, shipment, and load vs capacity view.",
            "Frozen-zone and firm-zone plan controls.",
            "Impact preview for order change, cancellation, capacity loss, and shipment pull-in.",
        ], [
            "Weekly plan becomes the coordination contract across departments.",
            "Plan cannot freeze while unresolved material, PCD, capacity, or wash conflicts remain.",
        ]),
        ("12. Daily Production Release Surface", "Release only what is actually ready and executable today.", [
            "Validate material, previous process, machine, manpower, quality, workcenter capacity, and next-process readiness.",
            "Check release impact against downstream wash capacity.",
            "Support release, block, exception approval, daily release history, and boundary-case holds.",
        ], [
            "Daily release becomes gate-based.",
            "Release does not create downstream wash overload without explicit warning.",
        ]),
        ("13. Workcenter Load Monitor", "Show load, capacity, queue, WIP ageing, utilization, and active constraints across workcenters.", [
            "Define capacity unit, planning bucket, primary and secondary constraint resource, normal capacity, overtime capacity, capacity loss triggers, and recovery levers.",
            "Monitor Fabric QC, Cutting, Sewing, Dry process, Wet wash, Drying, Finishing, Packing, Final QC, and Shipment documentation.",
        ], [
            "Bottleneck shifts are visible early.",
            "Recovery action is triggered before shipment risk becomes critical.",
        ]),
        ("14. Sewing Line Loading Surface", "Support realistic sewing line allocation and production target setting.", [
            "Assign orders to lines and calculate daily target from SMV, manpower, machines, skill, absenteeism, efficiency, learning curve, and quality performance.",
            "Distinguish gross target, quality-adjusted target, and net-good output.",
            "Handle absenteeism, machine breakdown, and style complexity changes as boundary cases.",
        ], [
            "Sewing plan reflects realistic net-good capacity.",
            "Line loading decisions improve resource utilization.",
        ]),
        ("15. Wash Planning Surface", "Plan wash/laundry as a denim value-creation, quality-risk, and capacity-constraining domain.", [
            "Define wash complexity class, fashion-effect requirement, dry process, wet wash route, repeat-cycle possibility, and rewash reason.",
            "Create wash batches, maintain shade-lot identity, track batch queue, track completion, reserve repeat-cycle capacity, and link outcome to finishing release.",
            "Capture cycle count, post-wash shade/effect QC outcome, capacity impact, and wash route performance analytics.",
        ], [
            "Rewash is visible as capacity-consuming work, not informal rework.",
            "Post-wash quality outcomes update WIP, capacity, and shipment risk.",
            "Wash queue overload triggers exception and recovery recommendation.",
        ]),
        ("16. WIP and Queue Monitoring Surface", "Monitor work waiting between processes.", [
            "Track WIP quantity, ageing, hold reason, next process, owner, shipment risk, threshold, and escalation trigger.",
            "Track shade-lot and batch identity where relevant.",
            "Distinguish normal queue from quality hold, rework hold, and capacity hold.",
        ], [
            "Hidden queues become visible.",
            "Ageing WIP escalates into the Exception and Alert Surface.",
        ]),
        ("17. Exception and Alert Surface", "Drive exception-based planning and recovery action.", [
            "Capture exception type, affected order, severity, owner, due date, suggested action, status, and escalation level.",
            "Add alerts for cancellation after cutting, capacity loss, overtime capacity addition, repeat wash beyond threshold, post-wash rejection, shipment pull-in, frozen plan change, and partial shipment.",
            "Use rule-based recommendations for recovery actions.",
        ], [
            "Critical issues are not buried in WhatsApp or Excel.",
            "Every critical issue has owner, due date, and visible action history.",
        ]),
        ("18. Shipment Readiness Surface", "Protect final shipment commitments.", [
            "Track shipment date, finished quantity, packed quantity, short quantity, final QC, AQL, cartons, labels, documentation, booking, and split shipment decision.",
            "Link shipment readiness to boundary-case events.",
            "Track premium freight or extra cost where applicable.",
        ], [
            "Production completion and dispatch readiness are clearly separated.",
            "Shipment risk is visible before dispatch date.",
        ]),
    ]
    for title, purpose, capabilities, success in surfaces:
        add_para(doc, title, style="Heading 2")
        add_para(doc, purpose)
        add_para(doc, "Key capabilities", style="Heading 3")
        add_bullets(doc, capabilities)
        add_para(doc, "Success criteria", style="Heading 3")
        add_bullets(doc, success)

    part(doc, "Part B: Mature-State Scope")
    add_para(doc, "The MVP stabilizes planning discipline. Mature deployment extends into deeper commercial, technical, execution, governance, and analytics capability.")
    add_table(doc, ["Surface", "Aligned Mature-State Intent"], [
        ("Executive Control Tower", "Management risk dashboard for shipment risk, bottlenecks, utilization, overtime exposure, recovery actions, and order-book health."),
        ("Enquiry and Costing", "Full costing and feasibility surface remains mature-state; MVP only captures milestone visibility unless business confirms immediate need."),
        ("Sampling and Approval Tracker", "Full sample workflow remains mature-state; MVP captures sample approval as lifecycle and PCD dependency."),
        ("Style Master and Technical File", "Maintain wash complexity class, fashion-effect requirement, route family, repeat-cycle probability, operation bulletin, and quality checkpoints."),
        ("BOM and Material Planning", "Material requirement calculation, fabric width and shrinkage impact, shortage alerts, and procurement linkage."),
        ("Procurement and Vendor Follow-up", "Vendor PO tracking, ETA, delay alerts, criticality ranking, and scorecard."),
        ("Fabric Inward and Fabric QC", "Roll-wise receipt, shade lot, 4-point inspection, width, GSM, shrinkage, skewing, stretch, and pass/fail/hold."),
        ("Operator Skill and Capacity", "Skill matrix, efficiency, quality rating, absenteeism, learning curve, and recommended allocation."),
        ("Cutting Room", "Cutting plan, shade-wise sequence, marker approval, spreading, cut quantity, bundle numbering, and issue to sewing."),
        ("Wash Recipe and Batch Execution", "Detailed recipe execution, machine assignment, parameters, shade/hand-feel/measurement result, and batch approval."),
        ("Quality Management", "QC across all stages with defect category, root cause, and quality-adjusted capacity feedback."),
        ("Rework and Recovery", "Rework order, affected quantity, responsible process, recovery time, capacity reservation, shipment impact, and closure."),
        ("Calendar / Gantt Planning", "Visual order, department, line, wash, shipment, approval, and procurement timelines."),
        ("Recovery Planning / What-If Simulation", "Scenario tests for fabric delay, line split, wash rework, overtime, shift extension, and shipment pull-forward."),
        ("Plan Change and Approval", "Frozen-zone and firm-zone plan change request, reason code, impact analysis, approval, and audit trail."),
        ("Department Handover", "Formal interdepartmental handovers with completeness checks."),
        ("Master Data Governance", "Completeness and version control for style, BOM, SMV, workcenter capacity, calendar, wash recipe, and vendor lead time."),
        ("Performance Analytics", "OTIF, utilization, plan adherence, line efficiency, net-good output, wash rework, quality defects, WIP ageing, and recovery cost."),
        ("Mobile / Shopfloor Update", "Low-friction output, defect, shortage, machine issue, absenteeism, and handover updates."),
        ("Role-Based Home Surfaces", "Role-specific dashboards for management, planners, merchandisers, procurement, QC, production, and shipment teams."),
    ], [2600, 6760], font_size=8.1)

    part(doc, "Part C: Functional Requirements")
    fr_rows = [
        ("FR-001", "System shall maintain order lifecycle status from confirmation to shipment."),
        ("FR-002", "System shall maintain planned and actual dates for each major order stage."),
        ("FR-003", "System shall enforce PCD readiness checks before cutting release."),
        ("FR-004", "System shall support weekly capacity-based production planning."),
        ("FR-005", "System shall support daily production release based on readiness."),
        ("FR-006", "System shall monitor load and capacity at each workcenter."),
        ("FR-007", "System shall identify current and future bottlenecks."),
        ("FR-008", "System shall support sewing line loading using SMV, manpower, and capacity assumptions."),
        ("FR-009", "System shall distinguish gross output from net-good output."),
        ("FR-010", "System shall support wash route and wash batch planning."),
        ("FR-011", "System shall track rewash and repeat wash capacity consumption."),
        ("FR-012", "System shall track WIP quantity and ageing between processes."),
        ("FR-013", "System shall generate exceptions for material, capacity, quality, approval, WIP, and shipment risks."),
        ("FR-014", "System shall assign owner and due date to every exception."),
        ("FR-015", "System shall track shipment readiness separately from production completion."),
        ("FR-016", "System shall maintain audit trail for plan changes and conditional releases."),
        ("FR-017", "System shall support role-based views and permissions."),
        ("FR-018", "System shall support integration with ERP or existing order/material data sources."),
        ("FR-019", "System shall support export and reporting for management review."),
        ("FR-020", "System shall preserve historical data for performance analytics."),
        ("FR-021", "System shall maintain product complexity class for denim and chino route planning."),
        ("FR-022", "System shall define MVP scheduling grain across order, style, color/shade-lot, stage, workcenter/line, date/shift, and wash batch."),
        ("FR-023", "System shall support rule-based impact preview for order cancellation, quantity change, shipment date change, and capacity change."),
        ("FR-024", "System shall treat wash route, dry/wet step, post-wash QC, and rewash cycle as first-class planning events."),
        ("FR-025", "System shall capture rewash reason, affected quantity, expected time, capacity impact, and approval status."),
        ("FR-026", "System shall maintain business-operational capacity definitions by workcenter."),
        ("FR-027", "System shall distinguish planner-controlled recommendations from automatic system changes."),
        ("FR-028", "System shall capture upstream milestones for enquiry, quotation, sampling, technical file receipt, and order confirmation."),
        ("FR-029", "System shall maintain audit trail for boundary-case events and scheduling reaction decisions."),
        ("FR-030", "System shall support seed scenarios or simulations for cancellation, capacity loss, rewash, and shipment pull-in before full planning-engine rollout."),
    ]
    add_table(doc, ["Req. ID", "Requirement"], fr_rows, [1150, 8210], font_size=8)

    part(doc, "Part D: Business Rules")
    rule_sections = [
        ("PCD Release Rules", [
            "No order should move to cutting without PCD readiness status.",
            "Blocked orders must have open dependency and owner.",
            "Conditional release must require reason, owner, expiry, and approval.",
            "Fabric QC failure should block PCD unless overridden by authorized role.",
        ]),
        ("Planning Rules", [
            "Weekly plan must be capacity checked before freeze.",
            "Frozen-zone changes must require reason code and approval.",
            "Daily release must validate real readiness.",
            "Workcenter overload should trigger exception.",
            "Orders with high shipment risk should receive priority review.",
        ]),
        ("Capacity Rules", [
            "Capacity should be based on agreed unit, planning bucket, available minutes, manpower, machine availability, SMV, efficiency, and demonstrated output.",
            "Net-good output should account for quality loss and rework.",
            "Rework and rewash must consume capacity.",
            "Wash capacity must be separately planned from sewing capacity.",
        ]),
        ("Scheduling Grain Rules", [
            "MVP scheduling shall operate at order x style x color/shade-lot x stage x workcenter/line x date/shift grain.",
            "Wash shall be planned at wash-batch grain.",
            "Operation-level data shall support SMV, line balance, and technical validation, but full operation-by-operation finite scheduling is outside MVP unless later validated.",
        ]),
        ("Boundary-Case Rules", [
            "Order cancellation, quantity change, delivery-date change, capacity change, rewash, QC failure, and shipment pull-in must create a visible system event.",
            "Each event must show impacted orders, affected capacity, shipment risk, suggested recovery, owner, and approval requirement.",
            "MVP should recommend and preview impact; it should not automatically reschedule frozen plans without approval.",
        ]),
        ("Rewash / Repeat Cycle Rules", [
            "Rewash must consume planned capacity.",
            "Rewash must carry reason, affected quantity, approval owner, expected duration, and shipment impact.",
            "Repeated wash beyond threshold must trigger exception.",
            "Post-wash QC result must control release to finishing.",
        ]),
        ("Automation Boundary Rules", [
            "MVP shall use rule-based impact preview and recommendations.",
            "Planner or authorized approver shall control plan changes in frozen or firm zones.",
            "Fully autonomous rescheduling is outside MVP unless validated by the business.",
        ]),
        ("Shipment Rules", [
            "Shipment readiness must include final QC, inspection, packing, documents, and booking.",
            "Production complete does not mean shipment ready.",
            "Short shipment risk must be visible before shipment date.",
            "Split shipment decision must be tracked and approved.",
        ]),
    ]
    for title, items in rule_sections:
        add_para(doc, title, style="Heading 2")
        add_bullets(doc, items)
    add_para(doc, "Order Cancellation Rules", style="Heading 2")
    add_table(doc, ["Cancellation Stage", "Required System Reaction"], [
        ("Before procurement", "Release planned capacity, cancel open procurement intent, and audit decision."),
        ("After procurement", "Flag material liability, review reallocation, and update capacity and cost exposure."),
        ("After cutting", "Freeze or reclassify cut WIP and require management decision."),
        ("During sewing", "Update WIP, capacity, and shipment plan; require planner action."),
        ("During/after wash", "Flag finished/semi-finished goods disposition and customer impact."),
    ], [2600, 6760])

    part(doc, "Part E: Non-Functional Requirements")
    add_table(doc, ["Category", "Requirement"], [
        ("Usability", "Screens must be simple, action-oriented, and suitable for planners and shopfloor supervisors."),
        ("Performance", "Workcenter load, WIP, and order-risk views should load quickly for large order volumes."),
        ("Scalability", "System must support large factory operations with many lines, workcenters, styles, and orders."),
        ("Reliability", "Planning data must be version-controlled and auditable."),
        ("Security", "Access should be role-based and protect sensitive commercial and customer data."),
        ("Auditability", "Plan changes, overrides, conditional releases, and shipment changes must be traceable."),
        ("Configurability", "Workcenters, statuses, thresholds, calendars, roles, product complexity, and capacity units must be configurable."),
        ("Explainability", "Planning recommendations must show why the system is recommending a recovery action."),
        ("Data Freshness", "Planning-critical screens must show data source and last update status."),
        ("Scenario Traceability", "Boundary-case outcomes must be traceable to the rule or user decision that caused them."),
        ("Change Governance", "Frozen-plan changes, overrides, and conditional releases must require controlled approval and audit."),
    ], [2200, 7160])

    part(doc, "Part F: Integration Requirements")
    add_table(doc, ["Integration Area", "Purpose"], [
        ("Order Management / ERP", "Order, customer, style, PO, quantity, delivery date."),
        ("BOM / Material System", "Fabric, trims, labels, packing materials, and consumables."),
        ("Procurement", "Vendor PO, ETA, inward status, delay alerts."),
        ("Sampling / Approval Records", "Upstream milestone visibility and PCD readiness dependency."),
        ("Fabric QC", "Inspection status, shade lot, shrinkage, and pass/fail/hold results."),
        ("Production Actuals", "Cutting, sewing, wash, finishing output."),
        ("Quality System", "Defects, holds, rework, AQL, post-wash QC."),
        ("HR / Attendance", "Operator availability and absenteeism."),
        ("Machine / Maintenance", "Machine availability, downtime, and breakdown."),
        ("Shipment / Logistics", "Packing, inspection, dispatch, documents, booking."),
        ("FastReact / Existing Planning Tool", "Define replacement, coexistence, import, or export role clearly."),
        ("Wash Route / Recipe Master", "Route family, standard cycle, repeat-cycle rules, and quality checkpoints."),
        ("Order Change / Cancellation Source", "Capture cancellation, quantity change, and delivery-date change events."),
        ("Capacity Calendar", "Shift, overtime, downtime, absenteeism, and approved capacity changes."),
    ], [2600, 6760])

    part(doc, "Part G: Roles and Users")
    add_table(doc, ["Role", "Key Usage"], [
        ("Management", "Control tower, shipment risk, bottlenecks, performance, approval escalation."),
        ("Production Planner", "Weekly plan, daily release, workcenter load, boundary-case recovery."),
        ("Merchandiser", "Order lifecycle, approvals, buyer comments, upstream milestones."),
        ("Procurement User", "Material status, vendor follow-up, shortage alerts."),
        ("Fabric QC User", "Fabric inspection, shade lot, shrinkage, and clearance."),
        ("Cutting Manager", "Cutting readiness, cut output, bundle handover."),
        ("Sewing Manager", "Line load, output, WIP, constraints, underperformance."),
        ("Washing Manager", "Wash plan, batch status, rewash, post-wash QC linkage."),
        ("Finishing Manager", "Finishing queue and packing readiness."),
        ("QC Manager", "Quality holds, inspection, defects, post-wash/final QC."),
        ("Shipment Team", "Shipment readiness, documents, booking, dispatch, split shipment."),
        ("Shopfloor Supervisor", "Daily updates, output, issues, handover."),
    ], [2500, 6860])

    part(doc, "Part H: Reporting Requirements")
    add_table(doc, ["Report", "Purpose"], [
        ("Order Status Report", "End-to-end order progress."),
        ("PCD Readiness Report", "Orders ready, blocked, conditional, or escalated."),
        ("Weekly Plan Report", "Department-wise weekly plan and freeze status."),
        ("Daily Release Report", "Work released, blocked, and conditionally released."),
        ("Workcenter Load Report", "Load vs capacity by workcenter."),
        ("Sewing Line Load Report", "Line-wise plan, target, output, and net-good result."),
        ("Wash Queue Report", "Wash load, queue, rewash, route, and post-wash QC."),
        ("WIP Ageing Report", "Stuck WIP, hold reason, owner, and shipment risk."),
        ("Exception Report", "Open issues and recovery actions."),
        ("Shipment Readiness Report", "Orders ready or at risk for shipment."),
        ("Scheduling Boundary Case Report", "Cancellation, capacity change, rewash, and shipment-pull events with system reaction."),
        ("Wash/Rewash Performance Report", "Wash queue, route performance, rewash count, reason, and capacity impact."),
        ("Capacity Change Log", "Approved and unapproved capacity changes by workcenter."),
        ("Order Change Impact Report", "Impact of cancellation, quantity change, and delivery-date change."),
        ("Planning Rulebook Compliance Report", "Whether critical releases followed defined scheduling rules."),
    ], [2700, 6660], font_size=8.2)

    part(doc, "Part I: Implementation Roadmap")
    add_para(doc, "Phase 0A: Scheduling Behaviour Rulebook and Simulation Validation", style="Heading 2")
    add_para(doc, "Before deep implementation, conduct a focused business workshop and produce the following deliverables:")
    add_numbered(doc, [
        "Capacity definition matrix.",
        "MVP planning grain decision.",
        "Order cancellation matrix.",
        "Capacity change matrix.",
        "Wash/rewash behaviour matrix.",
        "Shipment risk reaction matrix.",
        "Planner-control vs system-control boundary.",
        "Boundary-case seed scenario pack.",
        "BRD-to-build traceability matrix.",
    ])
    add_callout(doc, "Build gate", "Do not finalize daily release, workcenter load, or wash planning logic until Phase 0A confirms capacity definitions, wash/rework behaviour, and boundary-case reaction rules.")
    add_para(doc, "Phase 1: MVP Foundation", style="Heading 2")
    add_bullets(doc, ["Order Lifecycle", "PCD Readiness", "Weekly Planning Workbench", "Daily Production Release", "Workcenter Load Monitor", "Sewing Line Loading", "Wash Planning", "WIP and Queue Monitoring", "Exception and Alert Management", "Shipment Readiness"])
    add_para(doc, "Phase 2: Material, Sampling, and Technical Master Expansion", style="Heading 2")
    add_bullets(doc, ["Enquiry and costing", "Sampling and approvals", "Style technical file", "BOM and material planning", "Procurement and vendor follow-up", "Fabric inward and QC"])
    add_para(doc, "Phase 3: Shopfloor and Quality Depth", style="Heading 2")
    add_bullets(doc, ["Cutting room", "Wash recipe execution", "Quality management", "Rework and recovery", "Operator skill and capacity", "Department handover", "Mobile shopfloor updates"])
    add_para(doc, "Phase 4: Governance, Simulation, and Analytics", style="Heading 2")
    add_bullets(doc, ["Executive control tower", "Calendar / Gantt planning", "What-if simulation", "Plan change approval", "Master data governance", "Performance analytics", "Role-based home surfaces"])

    part(doc, "Part J: Risks and Mitigation")
    add_table(doc, ["Risk", "Mitigation"], [
        ("Users continue using Excel", "Make MVP surfaces action-critical and reduce duplicate entry."),
        ("Poor master data", "Start with minimum required master data and completeness checks."),
        ("Delayed shopfloor actuals", "Use simple update flows and supervisor-level capture."),
        ("Resistance to gate-based release", "Use management-backed policy for PCD and daily release."),
        ("FastReact overlap confusion", "Define integration or coexistence role clearly."),
        ("Wash complexity under-modeled", "Reframe wash as value-creation, quality-risk, and capacity domain; include in MVP."),
        ("Scheduling grain remains ambiguous", "Add MVP scheduling grain decision before build."),
        ("Boundary cases are discovered too late", "Add Scheduling Behaviour Rulebook and boundary-case test pack."),
        ("Premature automation creates user distrust", "Use rule-based recommendation and planner approval in MVP."),
        ("Pre-order scope expands MVP too far", "Capture upstream milestones in MVP; defer full costing unless confirmed."),
        ("Chinos over-modeled as denim", "Use product complexity classes and route configuration."),
        ("Too many alerts", "Prioritize shipment-impacting exceptions."),
    ], [3300, 6060], font_size=8.3)

    part(doc, "Part K: Open Decisions and Workshop Questions")
    question_groups = [
        ("Scope and Trigger", [
            "Was the original business trigger quotation/costing, production scheduling, delivery reliability, FastReact underuse, or all of these?",
            "Is the immediate pain before order confirmation, after order confirmation, or both?",
            "Are chinos handled in the same flow or a simpler production route?",
        ]),
        ("Current Planning", [
            "What planning is done in FastReact today?",
            "What planning is still done in Excel?",
            "Who owns the weekly plan and daily release?",
            "What is the current planning grain: order, PO, style/color, batch, line, operation, or shipment?",
        ]),
        ("Wash and Denim", [
            "What are the main wash route families?",
            "How is dry process planned separately from wet wash?",
            "How often does rewash occur?",
            "What are the main rewash reasons?",
            "Who approves rewash?",
            "Is drying or post-wash QC a separate bottleneck?",
        ]),
        ("Capacity", [
            "How is sewing capacity measured?",
            "How is wash capacity measured?",
            "How is absenteeism reflected?",
            "How is overtime approved?",
            "How are machine breakdowns captured?",
            "Which recovery levers are acceptable by workcenter?",
        ]),
        ("Boundary Cases", [
            "What happens when an order is cancelled before procurement, after fabric receipt, after cutting, or during wash?",
            "Should the system automatically reschedule or recommend actions?",
            "Who can approve frozen plan changes?",
        ]),
        ("Prototype Confidence", [
            "What minimum prototype will give the business confidence?",
            "Which real orders should be used as seed simulations?",
            "Which process should be simulated first: sewing, wash, or shipment risk?",
            "Which boundary cases must be proven before committed build?",
        ]),
    ]
    for title, qs in question_groups:
        add_para(doc, title, style="Heading 2")
        add_numbered(doc, qs)

    add_para(doc, "20. Final BRD Summary", style="Heading 1")
    add_para(doc, "Eratex requires an end-to-end planning and scheduling tool because its current planning environment appears to be partially digitized but still operationally dependent on Excel. The aligned BRD sharpens the root cause: denim manufacturing requires a governed operating spine that understands wash/laundry complexity, WIP movement, capacity changes, boundary cases, and exception ownership.")
    add_para(doc, "The proposed tool must become the daily operating discipline layer for Eratex. It must enforce order visibility, PCD readiness, weekly and daily planning control, workcenter load monitoring, realistic sewing line loading, denim-specific wash and rewash planning, WIP ageing visibility, exception ownership, shipment readiness, and controlled system reaction when factory reality changes.")
    add_callout(doc, "Final position", "Build a governed denim-aware operating spine, not just a scheduling screen.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
