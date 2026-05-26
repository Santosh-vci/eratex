from dataclasses import dataclass
from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "eratex_ui_visual_thesis_assets"
OUT_DOCX = ROOT / "Eratex_Planning_Scheduling_Tool_UI_Visual_Thesis_Handoff.docx"

BRAND = {
    "navy": "#183A5F",
    "blue": "#2E75B6",
    "blue_dark": "#1F4D78",
    "sky": "#E8F1F8",
    "ink": "#172033",
    "muted": "#657080",
    "line": "#D8E2EC",
    "bg": "#F5F7FA",
    "surface": "#FFFFFF",
    "green": "#2E7D32",
    "green_bg": "#E7F4EA",
    "amber": "#B7791F",
    "amber_bg": "#FFF4DB",
    "red": "#B42318",
    "red_bg": "#FCEBE8",
    "black": "#111827",
    "gray_bg": "#EEF2F6",
}


def hex_to_rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


F = {
    "xs": font(16),
    "sm": font(19),
    "base": font(22),
    "base_b": font(22, True),
    "md": font(25),
    "md_b": font(25, True),
    "lg": font(31, True),
    "xl": font(40, True),
}


def draw_text(draw, xy, text, fill=None, fnt=None, anchor=None):
    draw.text(xy, text, fill=fill or hex_to_rgb(BRAND["ink"]), font=fnt or F["base"], anchor=anchor)


def text_size(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw, text, fnt, max_width):
    words = str(text).split()
    lines = []
    current = ""
    for word in words:
        trial = word if not current else current + " " + word
        if text_size(draw, trial, fnt)[0] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_wrapped(draw, xy, text, max_width, fill=None, fnt=None, line_gap=7, max_lines=None):
    fnt = fnt or F["base"]
    x, y = xy
    lines = wrap_text(draw, text, fnt, max_width)
    if max_lines is not None and len(lines) > max_lines:
        lines = lines[:max_lines]
        while lines and text_size(draw, lines[-1] + "...", fnt)[0] > max_width:
            lines[-1] = lines[-1][:-1]
        lines[-1] = lines[-1] + "..."
    line_h = text_size(draw, "Ag", fnt)[1] + line_gap
    for idx, line in enumerate(lines):
        draw_text(draw, (x, y + idx * line_h), line, fill=fill, fnt=fnt)
    return y + len(lines) * line_h


def rounded(draw, xy, fill, outline=None, width=1, radius=6):
    draw.rounded_rectangle(xy, radius=radius, fill=hex_to_rgb(fill), outline=hex_to_rgb(outline) if outline else None, width=width)


def line(draw, xy, fill=None, width=1):
    draw.line(xy, fill=hex_to_rgb(fill or BRAND["line"]), width=width)


def pill(draw, x, y, text, fill, text_fill=None, outline=None, w=None):
    pad_x = 16
    h = 34
    tw, _ = text_size(draw, text, F["sm"])
    w = w or tw + pad_x * 2
    rounded(draw, (x, y, x + w, y + h), fill, outline=outline or fill, radius=6)
    draw_text(draw, (x + w / 2, y + h / 2 - 1), text, fill=hex_to_rgb(text_fill or BRAND["ink"]), fnt=F["sm"], anchor="mm")
    return x + w + 8


def status_color(status):
    return {
        "Ready": (BRAND["green_bg"], BRAND["green"]),
        "Green": (BRAND["green_bg"], BRAND["green"]),
        "Conditional": (BRAND["amber_bg"], BRAND["amber"]),
        "Yellow": (BRAND["amber_bg"], BRAND["amber"]),
        "Blocked": (BRAND["red_bg"], BRAND["red"]),
        "Red": (BRAND["red_bg"], BRAND["red"]),
        "Black": ("#E7E9ED", BRAND["black"]),
        "Open": (BRAND["amber_bg"], BRAND["amber"]),
        "Closed": (BRAND["green_bg"], BRAND["green"]),
    }.get(status, (BRAND["gray_bg"], BRAND["muted"]))


def draw_progress(draw, x, y, w, h, pct, color, bg="#E7ECF2", label=None):
    rounded(draw, (x, y, x + w, y + h), bg, radius=4)
    rounded(draw, (x, y, x + max(8, int(w * pct)), y + h), color, radius=4)
    if label:
        draw_text(draw, (x + w + 10, y + h / 2 - 1), label, fill=hex_to_rgb(BRAND["muted"]), fnt=F["sm"], anchor="lm")


def draw_table(draw, x, y, widths, row_h, headers, rows, header_fill=None, max_rows=None):
    header_fill = header_fill or BRAND["navy"]
    total_w = sum(widths)
    rounded(draw, (x, y, x + total_w, y + row_h), header_fill, radius=3)
    xx = x
    for idx, h in enumerate(headers):
        draw_wrapped(draw, (xx + 12, y + 9), h, widths[idx] - 20, fill=hex_to_rgb("#FFFFFF"), fnt=F["sm"], max_lines=1)
        xx += widths[idx]
    y += row_h
    visible_rows = rows[: max_rows or len(rows)]
    for ridx, row in enumerate(visible_rows):
        fill = "#FFFFFF" if ridx % 2 == 0 else "#F0F5F9"
        rounded(draw, (x, y, x + total_w, y + row_h), fill, outline=BRAND["line"], radius=1)
        xx = x
        for idx, cell in enumerate(row):
            draw_wrapped(draw, (xx + 12, y + 9), cell, widths[idx] - 20, fill=hex_to_rgb(BRAND["ink"]), fnt=F["sm"], max_lines=2)
            xx += widths[idx]
        y += row_h
    return y


def app_shell(title, active_nav, subtitle=None):
    img = Image.new("RGB", (1600, 980), hex_to_rgb(BRAND["bg"]))
    draw = ImageDraw.Draw(img)
    sidebar_w = 270
    top_h = 78
    rounded(draw, (0, 0, sidebar_w, 980), BRAND["navy"], radius=0)
    draw_text(draw, (30, 28), "ERATEX", fill="#FFFFFF", fnt=F["lg"])
    draw_text(draw, (30, 66), "Planning Control", fill="#B7D2EB", fnt=F["sm"])
    nav = [
        "Control Tower",
        "Order Lifecycle",
        "PCD Readiness",
        "Weekly Plan",
        "Daily Release",
        "Capacity",
        "Sewing Lines",
        "Wash Plan",
        "WIP Queues",
        "Exceptions",
        "Shipments",
    ]
    y = 132
    for item in nav:
        selected = item == active_nav
        fill = "#244C78" if selected else BRAND["navy"]
        rounded(draw, (18, y, sidebar_w - 18, y + 42), fill, radius=6)
        draw_text(draw, (36, y + 21), item, fill="#FFFFFF" if selected else "#B7C7D9", fnt=F["sm"], anchor="lm")
        y += 48
    line(draw, (sidebar_w, 0, sidebar_w, 980), fill="#C9D4E1")
    rounded(draw, (sidebar_w, 0, 1600, top_h), "#FFFFFF", radius=0)
    draw_text(draw, (300, 22), title, fnt=F["lg"], fill=hex_to_rgb(BRAND["ink"]))
    if subtitle:
        draw_text(draw, (302, 58), subtitle, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
    pill(draw, 1160, 22, "Factory: Eratex", "#EEF5FB", BRAND["blue_dark"], BRAND["line"], w=160)
    pill(draw, 1332, 22, "Week 22", "#EEF5FB", BRAND["blue_dark"], BRAND["line"], w=104)
    pill(draw, 1448, 22, "Live plan", BRAND["green_bg"], BRAND["green"], BRAND["green_bg"], w=106)
    content = (300, 110, 1550, 930)
    return img, draw, content


def panel(draw, xy, title=None, subtitle=None):
    rounded(draw, xy, "#FFFFFF", outline=BRAND["line"], radius=6)
    x1, y1, x2, _ = xy
    if title:
        draw_text(draw, (x1 + 20, y1 + 20), title, fnt=F["md_b"])
        if subtitle:
            draw_text(draw, (x1 + 20, y1 + 52), subtitle, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
        line(draw, (x1 + 20, y1 + 72, x2 - 20, y1 + 72))


def kpi(draw, x, y, w, label, value, trend, status="Green"):
    fill, color = status_color(status)
    rounded(draw, (x, y, x + w, y + 102), "#FFFFFF", outline=BRAND["line"], radius=6)
    draw_text(draw, (x + 18, y + 18), label, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
    draw_text(draw, (x + 18, y + 54), value, fnt=F["lg"], fill=hex_to_rgb(BRAND["ink"]))
    pill(draw, x + w - 106, y + 28, trend, fill, color, w=86)


def draw_heat_cells(draw, x, y, labels_y, labels_x, values, cell_w=116, cell_h=42):
    for i, label in enumerate(labels_x):
        draw_text(draw, (x + i * cell_w + cell_w / 2, y - 24), label, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]), anchor="mm")
    for r, row_label in enumerate(labels_y):
        draw_text(draw, (x - 12, y + r * cell_h + cell_h / 2), row_label, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]), anchor="rm")
        for c, v in enumerate(values[r]):
            if v >= 90:
                fill = BRAND["red_bg"]
                txt = BRAND["red"]
            elif v >= 78:
                fill = BRAND["amber_bg"]
                txt = BRAND["amber"]
            else:
                fill = BRAND["green_bg"]
                txt = BRAND["green"]
            rounded(draw, (x + c * cell_w, y + r * cell_h, x + (c + 1) * cell_w - 6, y + (r + 1) * cell_h - 6), fill, outline="#FFFFFF", radius=4)
            draw_text(draw, (x + c * cell_w + cell_w / 2 - 3, y + r * cell_h + cell_h / 2 - 3), f"{v}%", fnt=F["xs"], fill=hex_to_rgb(txt), anchor="mm")


def screen_control_tower(path):
    img, draw, c = app_shell("Executive Control Tower", "Control Tower", "Shipment risk, bottlenecks, utilization, and recovery ownership")
    x, y, w, h = c[0], c[1], c[2] - c[0], c[3] - c[1]
    kpi(draw, x, y, 280, "OTIF protected", "96.2%", "+1.1", "Green")
    kpi(draw, x + 306, y, 280, "Resource utilization", "72%", "+2.0", "Yellow")
    kpi(draw, x + 612, y, 280, "Red shipments", "18", "-4", "Red")
    kpi(draw, x + 918, y, 280, "Open recovery actions", "43", "-8", "Yellow")
    panel(draw, (x, y + 130, x + 740, y + 510), "Risk heatmap", "Workcenter load by week")
    labels_y = ["Fabric QC", "Cutting", "Sewing", "Wet wash", "Finishing", "Packing"]
    labels_x = ["W22", "W23", "W24", "W25", "W26"]
    values = [[62, 68, 75, 82, 71], [72, 74, 79, 88, 73], [92, 87, 76, 81, 80], [96, 91, 84, 77, 75], [73, 69, 64, 70, 73], [66, 72, 81, 86, 78]]
    draw_heat_cells(draw, x + 145, y + 235, labels_y, labels_x, values)
    panel(draw, (x + 770, y + 130, x + 1250, y + 510), "Top exception queue", "Management-relevant only")
    rows = [
        ("Fabric delayed", "PO-4481", "Procurement", "2d"),
        ("Wash over capacity", "ORD-1192", "Wash mgr", "1d"),
        ("AQL not scheduled", "ORD-1185", "QC mgr", "Today"),
        ("Line overload", "ORD-1202", "Planner", "3d"),
        ("Recovery overdue", "ORD-1177", "Sewing", "Today"),
    ]
    yy = y + 220
    for issue, order, owner, due in rows:
        rounded(draw, (x + 790, yy, x + 1230, yy + 48), "#FFFFFF", outline=BRAND["line"], radius=5)
        draw_text(draw, (x + 806, yy + 13), issue, fnt=F["sm"])
        draw_text(draw, (x + 1002, yy + 13), order, fnt=F["sm"], fill=hex_to_rgb(BRAND["blue_dark"]))
        draw_text(draw, (x + 1110, yy + 13), due, fnt=F["sm"], fill=hex_to_rgb(BRAND["red"]))
        yy += 56
    panel(draw, (x, y + 540, x + 1250, y + 820), "Order book health", "Customer-wise shipment risk and recovery exposure")
    draw_table(draw, x + 24, y + 620, [220, 150, 150, 170, 170, 250], 42, ["Customer", "Orders", "At risk", "OTIF exposure", "Utilization hit", "Required decision"], [
        ["Buyer A", "42", "7", "High", "Wet wash", "Approve overtime window"],
        ["Buyer B", "36", "3", "Medium", "Sewing L4", "Split line for cargo styles"],
        ["Buyer C", "29", "5", "High", "Fabric QC", "Escalate vendor ETA"],
        ["Buyer D", "21", "1", "Low", "Packing", "Hold dispatch review"],
    ])
    img.save(path)


def screen_order_lifecycle(path):
    img, draw, c = app_shell("Order Lifecycle Board", "Order Lifecycle", "One system status from confirmation to shipment")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 830, y + 805), "Order list with lifecycle stage", "Planner and merchandiser shared view")
    draw_table(draw, x + 22, y + 88, [110, 160, 125, 120, 120, 145, 110], 46, ["Order", "Buyer", "Style", "Qty", "Delivery", "Stage", "Risk"], [
        ["ORD-1185", "Northline", "DNM-421", "18,400", "Jun 12", "PCD", "Yellow"],
        ["ORD-1192", "Mode Co", "CHN-230", "12,200", "Jun 14", "Wash", "Red"],
        ["ORD-1202", "Kline", "JGR-008", "9,600", "Jun 18", "Sewing", "Green"],
        ["ORD-1210", "Vantage", "CRG-117", "14,800", "Jun 19", "Fabric QC", "Black"],
        ["ORD-1214", "Ridge", "DNM-712", "7,200", "Jun 22", "Finishing", "Yellow"],
        ["ORD-1220", "Everlane", "CHN-102", "20,000", "Jun 24", "Shipment", "Green"],
    ])
    panel(draw, (x + 865, y, x + 1250, y + 805), "Order detail drawer", "Visible without leaving the board")
    draw_text(draw, (x + 890, y + 92), "ORD-1192 - Mode Co", fnt=F["md_b"], fill=hex_to_rgb(BRAND["ink"]))
    x0, yy = x + 900, y + 150
    stages = ["Order", "Material", "PCD", "Cut", "Sew", "Wash", "Finish", "Ship"]
    colors = [BRAND["green"], BRAND["green"], BRAND["green"], BRAND["green"], BRAND["green"], BRAND["red"], BRAND["gray_bg"], BRAND["gray_bg"]]
    for i, st in enumerate(stages):
        cx = x0 + (i % 4) * 86
        cy = yy + (i // 4) * 92
        rounded(draw, (cx, cy, cx + 68, cy + 42), colors[i] if i < 6 else "#FFFFFF", outline=BRAND["line"], radius=5)
        draw_text(draw, (cx + 34, cy + 21), st, fnt=F["xs"], fill="#FFFFFF" if i < 6 else hex_to_rgb(BRAND["muted"]), anchor="mm")
        if i % 4 != 3:
            line(draw, (cx + 69, cy + 21, cx + 85, cy + 21), fill=BRAND["line"], width=2)
    draw_text(draw, (x + 890, y + 356), "Next action", fnt=F["md_b"])
    draw_wrapped(draw, (x + 890, y + 394), "Reserve additional wet-wash capacity or split batch across Machine W-03 and W-05.", 318, fnt=F["sm"])
    pill(draw, x + 890, y + 474, "Owner: Wash manager", "#EEF5FB", BRAND["blue_dark"], BRAND["line"], w=218)
    pill(draw, x + 890, y + 520, "Due: Tomorrow", BRAND["amber_bg"], BRAND["amber"], w=162)
    draw_text(draw, (x + 890, y + 598), "Activity", fnt=F["md_b"])
    for n, t in enumerate(["PCD released conditionally", "Sewing complete, 11,700 pcs", "Wash queue red - capacity breach"]):
        draw_text(draw, (x + 910, y + 640 + n * 38), t, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
    img.save(path)


def screen_pcd(path):
    img, draw, c = app_shell("PCD Readiness Gate", "PCD Readiness", "Release only orders with accountable readiness")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 142), "Decision strip", "PCD is a gate, not a status note")
    kpi(draw, x + 18, y + 24, 250, "Ready to cut", "28", "+6", "Green")
    kpi(draw, x + 292, y + 24, 250, "Conditional", "9", "+2", "Yellow")
    kpi(draw, x + 566, y + 24, 250, "Blocked", "14", "-3", "Red")
    kpi(draw, x + 840, y + 24, 250, "Escalated", "5", "-1", "Red")
    panel(draw, (x, y + 170, x + 790, y + 805), "Readiness checklist", "Every failed gate needs reason, owner, and due date")
    checks = [
        ("PO confirmed", "Ready"), ("BOM frozen", "Ready"), ("Fabric received", "Ready"), ("Fabric QC passed", "Blocked"),
        ("Shade lots mapped", "Ready"), ("Shrinkage report", "Conditional"), ("Trims secured", "Ready"), ("Pattern approved", "Ready"),
        ("Marker ready", "Ready"), ("PP sample approved", "Conditional"), ("Wash standard", "Ready"), ("Line allocated", "Blocked"),
        ("Wash capacity", "Blocked"), ("QC file ready", "Ready"),
    ]
    xx, yy = x + 26, y + 260
    for idx, (label, status) in enumerate(checks):
        col = idx % 2
        row = idx // 2
        bx = xx + col * 368
        by = yy + row * 62
        bg, fg = status_color(status)
        rounded(draw, (bx, by, bx + 340, by + 46), bg, outline=BRAND["line"], radius=5)
        draw_text(draw, (bx + 16, by + 23), label, fnt=F["sm"], anchor="lm")
        draw_text(draw, (bx + 326, by + 23), status, fnt=F["xs"], fill=hex_to_rgb(fg), anchor="rm")
    panel(draw, (x + 820, y + 170, x + 1250, y + 805), "Release decision", "Audit-friendly action surface")
    draw_text(draw, (x + 845, y + 252), "ORD-1210 - Cargo program", fnt=F["md_b"])
    draw_wrapped(draw, (x + 845, y + 300), "Blocked because Fabric QC failed and target line is not allocated. Conditional release requires management override.", 365, fnt=F["sm"])
    pill(draw, x + 845, y + 390, "Block release", BRAND["red_bg"], BRAND["red"], w=154)
    pill(draw, x + 1010, y + 390, "Request override", BRAND["amber_bg"], BRAND["amber"], w=180)
    draw_text(draw, (x + 845, y + 468), "Required fields for override", fnt=F["md_b"])
    for n, field in enumerate(["Reason code", "Approver", "Risk owner", "Expiry date", "Recovery action"]):
        rounded(draw, (x + 850, y + 510 + n * 48, x + 1210, y + 544 + n * 48), "#FFFFFF", outline=BRAND["line"], radius=3)
        draw_text(draw, (x + 864, y + 527 + n * 48), field, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]), anchor="lm")
    img.save(path)


def screen_weekly_plan(path):
    img, draw, c = app_shell("Weekly Planning Workbench", "Weekly Plan", "Capacity-checked 2-4 week firm plan with 8-12 week visibility")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Integrated weekly plan", "Order, material, cutting, sewing, wash, finishing, QC, and shipment in one view")
    days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
    start_x = x + 360
    for idx, d in enumerate(days):
        draw_text(draw, (start_x + idx * 130 + 55, y + 102), d, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]), anchor="mm")
        line(draw, (start_x + idx * 130, y + 122, start_x + idx * 130, y + 720), fill="#E3E9F0")
    rows = [
        ("ORD-1185 DNM-421", [("Material", 0, 1, BRAND["green"]), ("Cut", 1, 1, BRAND["blue"]), ("Sew L2", 2, 2, BRAND["blue_dark"]), ("Wash", 4, 1, BRAND["amber"])]),
        ("ORD-1192 CHN-230", [("Cut", 0, 1, BRAND["green"]), ("Sew L4", 1, 2, BRAND["blue_dark"]), ("Wet wash", 3, 2, BRAND["red"])]),
        ("ORD-1202 JGR-008", [("PCD", 0, 1, BRAND["amber"]), ("Cut", 2, 1, BRAND["blue"]), ("Sew L1", 3, 2, BRAND["blue_dark"])]),
        ("ORD-1214 DNM-712", [("Material", 0, 2, BRAND["red"]), ("PCD", 2, 1, BRAND["amber"]), ("Cut", 4, 1, BRAND["blue"])]),
        ("ORD-1220 CHN-102", [("Finish", 0, 1, BRAND["green"]), ("QC", 1, 1, BRAND["blue"]), ("Ship", 2, 1, BRAND["green"])]),
    ]
    yy = y + 145
    for name, bars in rows:
        draw_text(draw, (x + 36, yy + 24), name, fnt=F["sm"], fill=hex_to_rgb(BRAND["ink"]), anchor="lm")
        line(draw, (x + 28, yy + 58, x + 1228, yy + 58), fill="#EDF1F5")
        for label, start, span, color in bars:
            bx = start_x + start * 130 + 8
            by = yy + 8
            bw = span * 130 - 16
            rounded(draw, (bx, by, bx + bw, by + 34), color, radius=5)
            draw_text(draw, (bx + 10, by + 17), label, fnt=F["xs"], fill="#FFFFFF", anchor="lm")
        yy += 96
    panel(draw, (x + 30, y + 665, x + 1220, y + 780), "Freeze readiness")
    draw_text(draw, (x + 58, y + 742), "Capacity safe", fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
    draw_progress(draw, x + 58, y + 762, 300, 14, 0.76, BRAND["green"])
    draw_text(draw, (x + 372, y + 769), "76%", fnt=F["xs"], fill=hex_to_rgb(BRAND["green"]), anchor="lm")
    draw_text(draw, (x + 520, y + 742), "Wet wash overloaded", fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
    draw_progress(draw, x + 520, y + 762, 250, 14, 0.91, BRAND["red"])
    draw_text(draw, (x + 784, y + 769), "91%", fnt=F["xs"], fill=hex_to_rgb(BRAND["red"]), anchor="lm")
    pill(draw, x + 1032, y + 740, "Freeze blocked", BRAND["red_bg"], BRAND["red"], w=158)
    img.save(path)


def screen_daily_release(path):
    img, draw, c = app_shell("Daily Production Release", "Daily Release", "Only ready, executable work moves to the floor")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 390, y + 805), "Ready queue", "Released by workcenter")
    for idx, item in enumerate(["Cut: ORD-1185 - 4,200 pcs", "Sew L2: ORD-1202 - 1,100 pcs", "Wash W-03: ORD-1192 - 2 batches", "Finish: ORD-1220 - 8,000 pcs"]):
        rounded(draw, (x + 24, y + 95 + idx * 88, x + 360, y + 160 + idx * 88), BRAND["green_bg"], outline="#CFE5D3", radius=6)
        draw_wrapped(draw, (x + 42, y + 112 + idx * 88), item, 290, fnt=F["sm"])
    panel(draw, (x + 425, y, x + 815, y + 805), "Blocked queue", "Needs reason and owner")
    for idx, item in enumerate([("Fabric hold", "ORD-1210"), ("No line capacity", "ORD-1232"), ("QC file missing", "ORD-1238"), ("Machine down", "ORD-1240")]):
        rounded(draw, (x + 449, y + 95 + idx * 88, x + 785, y + 160 + idx * 88), BRAND["red_bg"], outline="#F4C8C1", radius=6)
        draw_text(draw, (x + 467, y + 116 + idx * 88), item[0], fnt=F["sm"])
        draw_text(draw, (x + 467, y + 140 + idx * 88), item[1], fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
    panel(draw, (x + 850, y, x + 1250, y + 805), "Release validation", "Gate-by-gate action panel")
    checks = [("Material", "Ready"), ("Previous process", "Ready"), ("Machine", "Ready"), ("Manpower", "Conditional"), ("Quality", "Ready"), ("Capacity", "Blocked"), ("Next process", "Ready")]
    yy = y + 94
    for label, st in checks:
        bg, fg = status_color(st)
        rounded(draw, (x + 878, yy, x + 1218, yy + 44), bg, outline=BRAND["line"], radius=5)
        draw_text(draw, (x + 896, yy + 22), label, fnt=F["sm"], anchor="lm")
        draw_text(draw, (x + 1200, yy + 22), st, fnt=F["xs"], fill=hex_to_rgb(fg), anchor="rm")
        yy += 54
    draw_text(draw, (x + 878, y + 526), "Actions", fnt=F["md_b"])
    pill(draw, x + 878, y + 570, "Release", BRAND["green_bg"], BRAND["green"], w=118)
    pill(draw, x + 1008, y + 570, "Block", BRAND["red_bg"], BRAND["red"], w=100)
    pill(draw, x + 878, y + 620, "Escalate", BRAND["amber_bg"], BRAND["amber"], w=130)
    img.save(path)


def screen_workcenter(path):
    img, draw, c = app_shell("Workcenter Load Monitor", "Capacity", "Bottleneck visibility across the end-to-end flow")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Load vs capacity heatmap", "Percent of available capacity by day")
    labels_y = ["Fabric QC", "Cutting", "Sewing L1", "Sewing L2", "Sewing L3", "Dry process", "Wet wash", "Drying", "Finishing", "Packing", "Final QC", "Shipment docs"]
    labels_x = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
    values = [
        [64, 72, 80, 69, 70, 61],
        [78, 82, 88, 73, 68, 65],
        [91, 86, 81, 77, 74, 70],
        [83, 79, 76, 82, 85, 88],
        [65, 69, 74, 71, 80, 84],
        [76, 81, 89, 94, 84, 74],
        [95, 97, 91, 88, 83, 79],
        [74, 80, 82, 79, 73, 72],
        [70, 68, 75, 81, 88, 84],
        [62, 66, 71, 79, 82, 74],
        [58, 60, 65, 70, 75, 81],
        [55, 62, 69, 76, 84, 88],
    ]
    draw_heat_cells(draw, x + 210, y + 150, labels_y, labels_x, values, cell_w=130, cell_h=48)
    panel(draw, (x + 935, y + 112, x + 1224, y + 690), "Constraint inspector", "Selected: Wet wash")
    draw_text(draw, (x + 958, y + 194), "Wet wash capacity breach", fnt=F["md_b"], fill=hex_to_rgb(BRAND["red"]))
    draw_wrapped(draw, (x + 958, y + 236), "Load exceeds safe threshold for three consecutive days. Reserve rewash capacity before accepting new batches.", 225, fnt=F["sm"])
    draw_progress(draw, x + 958, y + 350, 190, 18, 0.97, BRAND["red"], label="97% Tue")
    draw_progress(draw, x + 958, y + 400, 190, 18, 0.91, BRAND["red"], label="91% Wed")
    draw_progress(draw, x + 958, y + 450, 190, 18, 0.88, BRAND["amber"], label="88% Thu")
    pill(draw, x + 958, y + 536, "Create recovery", BRAND["amber_bg"], BRAND["amber"], w=180)
    img.save(path)


def screen_sewing(path):
    img, draw, c = app_shell("Sewing Line Loading", "Sewing Lines", "Realistic line allocation using SMV, manpower, and net-good output")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Line load board", "Plan, capacity, actuals, quality loss, and alternate-line suggestion")
    headers = ["Line", "Order / style", "SMV", "Manpower", "Target", "Load", "Plan vs actual", "Recommendation"]
    rows = [
        ["L1", "ORD-1202 / JGR-008", "12.4", "48", "1,850", "78%", "1,620 / 1,560", "Keep line"],
        ["L2", "ORD-1185 / DNM-421", "18.8", "52", "1,330", "96%", "1,280 / 1,090", "Review skill mix"],
        ["L3", "ORD-1214 / DNM-712", "16.2", "45", "1,210", "82%", "1,000 / 960", "Add helpers"],
        ["L4", "ORD-1192 / CHN-230", "10.7", "50", "2,050", "112%", "1,950 / 1,720", "Split to L5"],
        ["L5", "Available", "-", "46", "1,700", "44%", "-", "Accept split"],
    ]
    draw_table(draw, x + 24, y + 92, [68, 230, 80, 110, 110, 110, 170, 210], 54, headers, rows)
    panel(draw, (x + 24, y + 500, x + 620, y + 765), "Net-good output logic", "Capacity is not gross target")
    labels = [("Gross target", 0.92, BRAND["blue"]), ("Quality-adjusted target", 0.76, BRAND["amber"]), ("Actual net-good", 0.67, BRAND["red"])]
    yy = y + 590
    for label, pct, col in labels:
        draw_text(draw, (x + 56, yy + 8), label, fnt=F["sm"])
        draw_progress(draw, x + 250, yy, 250, 20, pct, col, label=f"{round(pct * 100)}%")
        yy += 54
    panel(draw, (x + 655, y + 500, x + 1226, y + 765), "Front-end behavior", "Inline edits should explain impact immediately")
    for n, t in enumerate(["Changing manpower recalculates target and load.", "Quality loss adjusts net-good output.", "Split-line action creates a plan-change request.", "Overload must create exception if not resolved."]):
        draw_wrapped(draw, (x + 686, y + 586 + n * 38), t, 485, fnt=F["sm"])
    img.save(path)


def screen_wash(path):
    img, draw, c = app_shell("Wash Planning Surface", "Wash Plan", "Wash is a core constraint, not a finishing afterthought")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Wash route and batch board", "Route, capacity, shade identity, rewash reserve, and finishing release linkage")
    lanes = [("Dry process", BRAND["blue"]), ("Wet wash", BRAND["blue_dark"]), ("Drying", BRAND["amber"]), ("Shade check", BRAND["green"]), ("Rewash loop", BRAND["red"])]
    lane_y = y + 125
    for lane_name, col in lanes:
        rounded(draw, (x + 30, lane_y, x + 1180, lane_y + 86), "#FFFFFF", outline=BRAND["line"], radius=5)
        draw_text(draw, (x + 54, lane_y + 22), lane_name, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
        for b in range(3):
            bx = x + 230 + b * 260 + (20 if lane_name == "Rewash loop" else 0)
            bw = 190 if b != 2 else 240
            rounded(draw, (bx, lane_y + 22, bx + bw, lane_y + 66), col, radius=5)
            draw_text(draw, (bx + 12, lane_y + 43), f"Batch {b + 1} - {['DNM', 'CHN', 'JGR'][b]}", fnt=F["xs"], fill="#FFFFFF", anchor="lm")
        lane_y += 105
    panel(draw, (x + 30, y + 670, x + 1180, y + 780), "Capacity guardrail", "Rewash reserve is explicit so urgent repeats do not silently consume the live plan")
    draw_progress(draw, x + 64, y + 735, 320, 18, 0.84, BRAND["amber"], label="Wet wash used")
    draw_progress(draw, x + 520, y + 735, 260, 18, 0.25, BRAND["red"], label="Rewash reserve")
    pill(draw, x + 918, y + 715, "Batch split needed", BRAND["amber_bg"], BRAND["amber"], w=190)
    img.save(path)


def screen_wip(path):
    img, draw, c = app_shell("WIP and Queue Monitoring", "WIP Queues", "Hidden waiting time becomes visible and owned")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Process swimlanes", "WIP cards carry quantity, age, owner, next process, and shipment risk")
    stages = ["Fabric QC", "Cut panels", "Sewing out", "Wash queue", "Finishing", "Packing", "Inspection"]
    sx = x + 35
    for idx, st in enumerate(stages):
        lx = sx + idx * 170
        rounded(draw, (lx, y + 100, lx + 150, y + 690), "#F8FAFC", outline=BRAND["line"], radius=5)
        draw_text(draw, (lx + 75, y + 128), st, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]), anchor="mm")
    cards = [
        (0, 0, "ORD-1210", "6,400", "18h", "Yellow"),
        (1, 1, "ORD-1185", "3,200", "9h", "Green"),
        (2, 0, "ORD-1202", "1,800", "6h", "Green"),
        (3, 0, "ORD-1192", "11,700", "42h", "Red"),
        (3, 2, "ORD-1231", "2,100", "29h", "Yellow"),
        (4, 1, "ORD-1214", "5,600", "14h", "Yellow"),
        (5, 0, "ORD-1220", "8,000", "7h", "Green"),
        (6, 1, "ORD-1177", "3,400", "31h", "Red"),
    ]
    for stage, row, order, qty, age, risk in cards:
        lx = sx + stage * 170 + 12
        cy = y + 165 + row * 138
        bg, fg = status_color(risk)
        rounded(draw, (lx, cy, lx + 126, cy + 104), bg, outline=BRAND["line"], radius=6)
        draw_text(draw, (lx + 12, cy + 16), order, fnt=F["xs"], fill=hex_to_rgb(fg))
        draw_text(draw, (lx + 12, cy + 45), qty, fnt=F["md_b"], fill=hex_to_rgb(BRAND["ink"]))
        draw_text(draw, (lx + 12, cy + 76), f"Age {age}", fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
    panel(draw, (x + 35, y + 710, x + 1188, y + 780), "Escalation rule", "If age exceeds threshold, create exception with owner and due date; do not rely on visual color alone")
    img.save(path)


def screen_exceptions(path):
    img, draw, c = app_shell("Exception and Alert Center", "Exceptions", "Exception-based planning with owner, due date, status, and recovery action")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 850, y + 805), "Exception board", "Grouped by severity and planning impact")
    cols = [("Red", BRAND["red_bg"], BRAND["red"]), ("Yellow", BRAND["amber_bg"], BRAND["amber"]), ("Green", BRAND["green_bg"], BRAND["green"])]
    for idx, (name, bg, fg) in enumerate(cols):
        cx = x + 28 + idx * 268
        rounded(draw, (cx, y + 94, cx + 246, y + 720), "#FFFFFF", outline=BRAND["line"], radius=6)
        draw_text(draw, (cx + 20, y + 124), name, fnt=F["md_b"], fill=hex_to_rgb(fg))
    issues = [
        (0, "Wash queue exceeds capacity", "ORD-1192", "Owner: Wash mgr"),
        (0, "Shipment risk turned red", "ORD-1177", "Owner: Shipment"),
        (1, "PP sample approval pending", "ORD-1214", "Owner: Merch"),
        (1, "Line overloaded next week", "ORD-1202", "Owner: Planner"),
        (2, "Recovery action closed", "ORD-1185", "Owner: Sewing"),
    ]
    offset = [0, 0, 0]
    for col, title, order, owner in issues:
        cx = x + 48 + col * 268
        cy = y + 170 + offset[col] * 145
        bg, fg = status_color(["Red", "Yellow", "Green"][col])
        rounded(draw, (cx, cy, cx + 206, cy + 116), bg, outline=BRAND["line"], radius=6)
        draw_wrapped(draw, (cx + 14, cy + 14), title, 178, fnt=F["sm"], max_lines=2)
        draw_text(draw, (cx + 14, cy + 70), order, fnt=F["xs"], fill=hex_to_rgb(BRAND["blue_dark"]))
        draw_text(draw, (cx + 14, cy + 92), owner, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
        offset[col] += 1
    panel(draw, (x + 885, y, x + 1250, y + 805), "Exception detail", "Action log and resolution workflow")
    draw_text(draw, (x + 910, y + 96), "Wash queue exceeds capacity", fnt=F["md_b"], fill=hex_to_rgb(BRAND["red"]))
    for idx, field in enumerate([("Affected order", "ORD-1192"), ("Severity", "Red"), ("Owner", "Wash manager"), ("Due date", "Today"), ("Status", "Open")]):
        draw_text(draw, (x + 910, y + 160 + idx * 42), field[0], fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]))
        draw_text(draw, (x + 1055, y + 160 + idx * 42), field[1], fnt=F["sm"])
    draw_text(draw, (x + 910, y + 410), "Suggested recovery", fnt=F["md_b"])
    draw_wrapped(draw, (x + 910, y + 452), "Split CHN-230 wash batch and reserve W-05 for rewash load. Notify finishing of revised release.", 296, fnt=F["sm"])
    pill(draw, x + 910, y + 570, "Accept action", BRAND["green_bg"], BRAND["green"], w=150)
    pill(draw, x + 910, y + 620, "Request change", BRAND["amber_bg"], BRAND["amber"], w=174)
    img.save(path)


def screen_shipment(path):
    img, draw, c = app_shell("Shipment Readiness", "Shipments", "Production complete is not dispatch ready")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Readiness matrix", "Final QC, inspection, packing, documents, and booking separated from production completion")
    headers = ["Order", "Final QC", "AQL", "Packing", "Labels", "Docs", "Booking", "Risk", "Action"]
    rows = [
        ["ORD-1177", "Pass", "Pending", "92%", "Ready", "Missing", "Booked", "Red", "Escalate docs"],
        ["ORD-1185", "Pass", "Pass", "100%", "Ready", "Ready", "Booked", "Green", "Release"],
        ["ORD-1192", "Hold", "Not set", "68%", "Ready", "Draft", "Open", "Red", "Recover wash"],
        ["ORD-1214", "Pass", "Pending", "88%", "Mismatch", "Ready", "Open", "Yellow", "Fix labels"],
        ["ORD-1220", "Pass", "Pass", "100%", "Ready", "Ready", "Booked", "Green", "Dispatch"],
    ]
    draw_table(draw, x + 24, y + 94, [110, 110, 95, 110, 110, 110, 110, 90, 180], 52, headers, rows)
    panel(draw, (x + 24, y + 460, x + 590, y + 765), "Dispatch calendar", "Visible short shipment and split shipment decisions")
    for i, day in enumerate(["Jun 10", "Jun 11", "Jun 12", "Jun 13", "Jun 14"]):
        rounded(draw, (x + 48 + i * 104, y + 540, x + 134 + i * 104, y + 704), "#FFFFFF", outline=BRAND["line"], radius=5)
        draw_text(draw, (x + 91 + i * 104, y + 568), day, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]), anchor="mm")
        if i in [2, 4]:
            bg, fg = status_color("Red" if i == 4 else "Yellow")
            rounded(draw, (x + 60 + i * 104, y + 600, x + 122 + i * 104, y + 658), bg, radius=5)
            draw_text(draw, (x + 91 + i * 104, y + 628), "Risk", fnt=F["xs"], fill=hex_to_rgb(fg), anchor="mm")
    panel(draw, (x + 630, y + 460, x + 1226, y + 765), "UX rule", "Shipment team must see blockers before dispatch day")
    for n, t in enumerate(["Separate production completion from dispatch readiness.", "Make missing documents and inspection bookings first-class blockers.", "Support split shipment decision and approval trail.", "Expose short quantity before final packing handover."]):
        draw_wrapped(draw, (x + 660, y + 544 + n * 42), t, 510, fnt=F["sm"])
    img.save(path)


def screen_shopfloor(path):
    img, draw, c = app_shell("Shopfloor Update Surface", "Daily Release", "Low-friction mobile and tablet update path")
    x, y = c[0], c[1]
    panel(draw, (x, y, x + 1250, y + 805), "Supervisor tablet pattern", "Fast capture for output, issues, defects, and handover confirmation")
    phone_x = x + 92
    rounded(draw, (phone_x, y + 78, phone_x + 370, y + 740), "#111827", radius=28)
    rounded(draw, (phone_x + 20, y + 112, phone_x + 350, y + 710), "#F8FAFC", radius=18)
    draw_text(draw, (phone_x + 44, y + 148), "Line L4 update", fnt=F["md_b"])
    draw_text(draw, (phone_x + 44, y + 184), "ORD-1192 / CHN-230", fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
    for idx, (label, value) in enumerate([("Output", "1,720 pcs"), ("Defects", "86 pcs"), ("Shortage", "None"), ("Machine issue", "Resolved")]):
        rounded(draw, (phone_x + 44, y + 238 + idx * 70, phone_x + 326, y + 292 + idx * 70), "#FFFFFF", outline=BRAND["line"], radius=6)
        draw_text(draw, (phone_x + 62, y + 258 + idx * 70), label, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
        draw_text(draw, (phone_x + 304, y + 258 + idx * 70), value, fnt=F["sm"], anchor="rm")
    pill(draw, phone_x + 62, y + 566, "Submit update", BRAND["blue"], "#FFFFFF", w=236)
    panel(draw, (x + 545, y + 90, x + 1188, y + 378), "Design requirements", "Built for factory floor use")
    for n, t in enumerate(["Large tap targets and short forms.", "Offline-friendly queue for weak network areas.", "Role-scoped tasks instead of full planning views.", "Every update writes back to plan, WIP, and exception records."]):
        draw_wrapped(draw, (x + 574, y + 162 + n * 44), t, 558, fnt=F["sm"])
    panel(draw, (x + 545, y + 420, x + 1188, y + 740), "State model", "Do not block supervisor capture when sync fails")
    draw_table(draw, x + 574, y + 510, [170, 360], 48, ["State", "Front-end behavior"], [
        ["Draft", "Validate required fields locally"],
        ["Queued", "Show pending sync badge"],
        ["Synced", "Lock audit fields; allow correction workflow"],
        ["Rejected", "Display reason and retry path"],
    ])
    img.save(path)


def screen_system_map(path):
    img = Image.new("RGB", (1600, 920), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw_text(draw, (70, 55), "UI operating model", fnt=F["xl"], fill=hex_to_rgb(BRAND["ink"]))
    draw_text(draw, (72, 105), "The front end is organized around flow, readiness, capacity, exception ownership, and shipment protection.", fnt=F["md"], fill=hex_to_rgb(BRAND["muted"]))
    stages = [
        ("Order", "Lifecycle board"),
        ("Readiness", "PCD gate"),
        ("Plan", "Weekly workbench"),
        ("Release", "Daily console"),
        ("Execute", "Line and wash"),
        ("Monitor", "WIP queues"),
        ("Recover", "Exceptions"),
        ("Ship", "Readiness"),
    ]
    x0, y0 = 80, 240
    for idx, (top, bottom) in enumerate(stages):
        x = x0 + idx * 180
        rounded(draw, (x, y0, x + 150, y0 + 110), BRAND["sky"], outline=BRAND["line"], radius=8)
        draw_text(draw, (x + 75, y0 + 38), top, fnt=F["md_b"], fill=hex_to_rgb(BRAND["blue_dark"]), anchor="mm")
        draw_text(draw, (x + 75, y0 + 74), bottom, fnt=F["xs"], fill=hex_to_rgb(BRAND["muted"]), anchor="mm")
        if idx < len(stages) - 1:
            line(draw, (x + 152, y0 + 55, x + 178, y0 + 55), fill=BRAND["blue"], width=3)
    roles = [
        ("Management", "Control tower and risk review"),
        ("Planner", "Plan, capacity, release, recovery"),
        ("Merchandiser", "Order lifecycle and approvals"),
        ("Factory managers", "Line, wash, WIP, output"),
        ("Shipment team", "Dispatch readiness"),
    ]
    y = 460
    draw_text(draw, (80, y), "Role-specific home surfaces", fnt=F["lg"], fill=hex_to_rgb(BRAND["ink"]))
    for idx, (role, focus) in enumerate(roles):
        bx = 80 + (idx % 3) * 480
        by = y + 70 + (idx // 3) * 115
        rounded(draw, (bx, by, bx + 430, by + 80), "#F8FAFC", outline=BRAND["line"], radius=6)
        draw_text(draw, (bx + 20, by + 22), role, fnt=F["md_b"], fill=hex_to_rgb(BRAND["blue_dark"]))
        draw_text(draw, (bx + 20, by + 54), focus, fnt=F["sm"], fill=hex_to_rgb(BRAND["muted"]))
    img.save(path)


@dataclass
class Screen:
    title: str
    image: str
    primary_users: str
    purpose: str
    data_inputs: str
    core_interactions: str
    states: str
    dev_notes: str


SCREENS = [
    Screen("Executive Control Tower", "screen_01_control_tower.png", "Management, head planner", "Show factory-wide risk, bottlenecks, OTIF exposure, and recovery ownership.", "Order book, capacity, WIP, exceptions, shipment readiness, recovery cost.", "Filter by week/customer/workcenter; drill into exception; create or approve recovery action.", "Fresh, stale, no-risk, high-risk, integration unavailable.", "Use aggregate APIs with cached summaries; preserve drill-down links to the underlying order and exception records."),
    Screen("Order Lifecycle Board", "screen_02_order_lifecycle.png", "Planner, merchandiser, production planner", "Give every order a single lifecycle status and next action.", "Order master, planned/actual dates, stage owner, dependencies, risk, comments.", "Search, filter, open detail drawer, update next action, hand off to linked surfaces.", "No orders, partial ERP sync, delayed stage, blocked stage, archived order.", "Board must support high row counts through virtualization and stable row height."),
    Screen("PCD Readiness Gate", "screen_03_pcd.png", "Planner, cutting manager, fabric QC, merchandiser", "Prevent cutting release without visible readiness and owner accountability.", "PO, BOM, fabric QC, trims, pattern, marker, PP sample, wash standard, line allocation.", "Approve ready release, block release, request conditional release, assign blocker owner.", "Ready, conditional, blocked, escalated, override expired.", "Every conditional release needs audit trail fields and immutable approval metadata."),
    Screen("Weekly Planning Workbench", "screen_04_weekly_plan.png", "Production planner, factory managers", "Convert order book into a capacity-checked weekly coordination contract.", "Orders, workcenter calendars, capacity, material readiness, shipment dates, line load.", "Drag schedule bars, freeze plan, inspect impact, resolve conflicts, compare versions.", "Draft, frozen, change pending, conflict, overload, stale capacity.", "Use a timeline component with sticky rows, left frozen columns, and conflict overlays."),
    Screen("Daily Production Release", "screen_05_daily_release.png", "Planner, department managers, shopfloor supervisors", "Release only work that is ready and executable today.", "Weekly plan, readiness gates, previous process completion, manpower, machine availability, QC clearance.", "Release, block, escalate, add reason code, view release history.", "Ready queue, blocked queue, conditional approval, offline update, duplicate release.", "All action buttons must be disabled until validations and required reason fields are satisfied."),
    Screen("Workcenter Load Monitor", "screen_06_workcenter.png", "Planner, production managers, management", "Expose shifting bottlenecks before they become shipment risk.", "Capacity by workcenter, planned load, WIP, queue age, throughput, recovery flags.", "Change date range, inspect cell, create recovery, compare available vs planned.", "Normal, warning, overload, missing capacity calendar, stale actuals.", "Heatmap cells need keyboard focus, tooltips, and click-to-filter behavior."),
    Screen("Sewing Line Loading", "screen_07_sewing.png", "Sewing manager, planner", "Allocate orders to sewing lines using realistic net-good capacity.", "SMV, manpower, machine availability, skill level, absenteeism, efficiency, quality loss.", "Assign line, split order, edit manpower, compare plan vs actual, create line overload exception.", "Underloaded, overloaded, actuals delayed, quality loss high, skill mismatch.", "Calculations should be shared between UI and backend or returned by a planning-calculation API."),
    Screen("Wash Planning Surface", "screen_08_wash.png", "Washing manager, planner", "Plan dry process, wet wash, rewash, and batch capacity as hard constraints.", "Wash route, batch size, machine capacity, shade lot, recipe, rewash requirement, finishing release.", "Create batch, sequence route, reserve rewash capacity, approve release to finishing.", "Batch queued, in process, hold, rewash, passed, capacity breach.", "Route visualization must preserve shade-lot identity and show rewash as capacity-consuming work."),
    Screen("WIP and Queue Monitoring", "screen_09_wip.png", "Planner, department managers, supervisors", "Surface hidden waiting time between processes.", "Stage inventory, quantity, queue age, owner, next process, hold reason, shipment risk.", "Filter by age, owner, stage, shipment risk; escalate WIP; confirm handover.", "Normal age, threshold breach, quality hold, missing handover, owner unassigned.", "Use lane virtualization if the board grows; card density should remain stable across zoom levels."),
    Screen("Exception and Alert Center", "screen_10_exceptions.png", "Planner, management, functional owners", "Turn informal issues into structured recovery work.", "Exception type, affected order, severity, owner, due date, suggested action, status, escalation level.", "Accept action, request change, assign owner, close, escalate, view history.", "Open, in progress, blocked, overdue, closed, reopened.", "Alerts must be deduplicated; same issue should not spawn multiple competing tickets."),
    Screen("Shipment Readiness", "screen_11_shipment.png", "Shipment team, QC manager, planner", "Separate production completion from dispatch readiness.", "Final QC, AQL, packing, barcode/label, documents, forwarder booking, short quantity.", "Confirm readiness, escalate blocker, approve split shipment, release dispatch.", "Ready, missing docs, AQL pending, label mismatch, booking open, short shipment.", "Completion percent is not enough; show the exact blocker and responsible function."),
    Screen("Shopfloor Update Surface", "screen_12_shopfloor.png", "Shopfloor supervisor, department managers", "Capture actuals and issues with minimum friction from tablet or mobile.", "Released work, output, defects, shortages, machine issues, handover confirmations.", "Submit update, queue offline, attach issue, confirm handover, correct rejected update.", "Draft, queued, synced, rejected, duplicate, offline.", "Mobile should be task-scoped and resilient to sync failure; avoid exposing full planning controls."),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_border(cell, color="D8E2EC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(width_dxa))
    tblInd = tblPr.first_child_found_in("w:tblInd")
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:type"), "dxa")
    tblInd.set(qn("w:w"), str(indent_dxa))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def paragraph_border_bottom(paragraph, color="D8E2EC", size="8"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(23, 32, 51)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.text = "Eratex Planning & Scheduling Tool | Front-End UI Visual Thesis"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(101, 112, 128)
    paragraph_border_bottom(p, "D8E2EC", "4")

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Confidential - Eratex Internal Use Only"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(101, 112, 128)
    p2 = footer.add_paragraph()
    add_page_number(p2)
    for r in p2.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(101, 112, 128)


def add_para(doc, text="", style=None, bold=False, color=None, size=None, after=None, before=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    if size:
        r.font.size = Pt(size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)


def add_label_table(doc, rows, widths=None, header=False):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    if widths is None:
        widths = [Inches(2.1), Inches(4.2)]
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            cell.width = widths[idx] if idx < len(widths) else Inches(1)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if ridx == 0 and header:
                set_cell_shading(cell, "E8EEF5")
                run = cell.paragraphs[0].add_run(text)
                run.bold = True
            else:
                run = cell.paragraphs[0].add_run(text)
                if idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("1F4D78")
            run.font.size = Pt(9.5)
    return table


def add_screen_spec_table(doc, screen):
    rows = [
        ("Primary users", screen.primary_users),
        ("Purpose", screen.purpose),
        ("Core data inputs", screen.data_inputs),
        ("Primary interactions", screen.core_interactions),
        ("Required states", screen.states),
        ("Front-end handoff notes", screen.dev_notes),
    ]
    add_label_table(doc, rows, widths=[Inches(1.55), Inches(4.85)])


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F9")
    set_cell_border(cell, "D8E2EC")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("1F4D78")
    p.add_run(text)
    return table


def add_mockup(doc, path, caption):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.45))
    cap = add_para(doc, caption, size=8.5, color="657080", after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.keep_with_next = False


def generate_assets():
    ASSET_DIR.mkdir(exist_ok=True)
    generators = [
        ("screen_00_system_map.png", screen_system_map),
        ("screen_01_control_tower.png", screen_control_tower),
        ("screen_02_order_lifecycle.png", screen_order_lifecycle),
        ("screen_03_pcd.png", screen_pcd),
        ("screen_04_weekly_plan.png", screen_weekly_plan),
        ("screen_05_daily_release.png", screen_daily_release),
        ("screen_06_workcenter.png", screen_workcenter),
        ("screen_07_sewing.png", screen_sewing),
        ("screen_08_wash.png", screen_wash),
        ("screen_09_wip.png", screen_wip),
        ("screen_10_exceptions.png", screen_exceptions),
        ("screen_11_shipment.png", screen_shipment),
        ("screen_12_shopfloor.png", screen_shopfloor),
    ]
    for filename, maker in generators:
        maker(ASSET_DIR / filename)


def build_docx():
    generate_assets()
    doc = Document()
    configure_doc(doc)

    add_para(doc, "Front-End UI Visual Thesis", size=12, bold=True, color="2E75B6", after=2)
    add_para(doc, "Eratex Planning & Scheduling Tool", size=28, bold=True, color="183A5F", after=4)
    add_para(doc, "Handoff document for UI engineering, product design, and implementation planning.", size=13, color="657080", after=18)
    add_label_table(
        doc,
        [
            ("Company", "Eratex"),
            ("Source", "Eratex Planning Scheduling Tool BRD, Version 1.0"),
            ("Document purpose", "Translate BRD surfaces into front-end information architecture, UI patterns, mockups, states, and build priorities."),
            ("Audience", "Front-end UI development team, product owner, UX designer, QA lead"),
            ("Date", "2026-05-26"),
            ("Status", "Professional handoff draft"),
        ],
        widths=[Inches(1.75), Inches(4.65)],
    )
    add_para(doc, "", after=8)
    add_mockup(doc, ASSET_DIR / "screen_00_system_map.png", "Figure 1 - UI operating model derived from the BRD.")
    doc.add_page_break()

    add_para(doc, "1. Visual Thesis", style="Heading 1")
    add_callout(
        doc,
        "Visual thesis",
        "The product should feel like a factory control room: dense, calm, precise, and action-oriented, with risk signals visible without turning the interface into an alarm wall.",
    )
    add_para(doc, "Interaction thesis", style="Heading 2")
    add_bullets(
        doc,
        [
            "Users should move from risk detection to owner assignment to recovery action without leaving the current screen.",
            "Planning tables and timelines should keep context stable through frozen columns, sticky headers, and detail drawers.",
            "Color must communicate status, but every risk signal also needs text, owner, due date, and next action.",
        ]
    )
    add_para(doc, "Front-end design posture", style="Heading 2")
    add_bullets(
        doc,
        [
            "Use a restrained operational UI: left navigation, persistent filters, dense data tables, heatmaps, timelines, and drawers.",
            "Avoid marketing-style dashboards. The first screen for each role should expose actual work, blockers, and decisions.",
            "Use compact panels only when they frame an interaction, not as decoration. Data density is acceptable when hierarchy is clear.",
            "Treat exceptions, readiness gates, and plan changes as first-class workflows, not as comments on records.",
        ]
    )

    add_para(doc, "2. Information Architecture", style="Heading 1")
    add_para(
        doc,
        "The application shell should organize the planning workflow around the factory flow: order visibility, readiness, weekly planning, daily release, capacity, execution, WIP, exceptions, and shipment readiness.",
    )
    add_label_table(
        doc,
        [
            ("Navigation group", "Representative screens"),
            ("Control", "Executive Control Tower, role-based home surfaces, business success metrics"),
            ("Order visibility", "Order Lifecycle Board, customer/order detail drawer, stage history"),
            ("Readiness", "PCD Readiness Gate, blocker ownership, conditional release approval"),
            ("Planning", "Weekly Planning Workbench, Gantt planning, plan change approval"),
            ("Execution release", "Daily Production Release, shopfloor update, department handover"),
            ("Capacity", "Workcenter Load Monitor, Sewing Line Loading, Wash Planning"),
            ("Flow monitoring", "WIP and Queue Monitoring, exception creation from aged WIP"),
            ("Recovery", "Exception and Alert Center, rework and recovery surface, what-if simulation"),
            ("Dispatch", "Shipment Readiness, shipment calendar, split shipment approval"),
            ("Data foundation", "Style master, BOM, vendor follow-up, fabric QC, master data governance"),
        ],
        widths=[Inches(1.65), Inches(4.75)],
        header=True,
    )
    add_para(doc, "Role-to-screen ownership", style="Heading 2")
    add_label_table(
        doc,
        [
            ("Role", "Primary surfaces"),
            ("Management", "Control Tower, exceptions, shipment risk, utilization, performance analytics"),
            ("Production planner", "Weekly plan, daily release, workcenter load, recovery planning, plan changes"),
            ("Merchandiser", "Order lifecycle, sampling approvals, buyer comments, shipment risk context"),
            ("Procurement", "Material status, vendor ETA, shortage alerts, PCD blockers"),
            ("Fabric QC", "Fabric inward, inspection, shade lot, pass/fail/hold clearance"),
            ("Cutting manager", "PCD gate, cutting plan, ready-to-cut release, cut output"),
            ("Sewing manager", "Line loading, daily output, line WIP, underperformance"),
            ("Washing manager", "Wash route, batch board, capacity, rewash reserve"),
            ("Shipment team", "Shipment readiness, documents, booking, split shipment decisions"),
            ("Shopfloor supervisor", "Daily task update, output, defects, shortages, handovers"),
        ],
        widths=[Inches(1.6), Inches(4.8)],
        header=True,
    )
    doc.add_page_break()

    add_para(doc, "3. Product UI System", style="Heading 1")
    add_para(doc, "Shell and layout", style="Heading 2")
    add_bullets(
        doc,
        [
            "Use a persistent left navigation with module names matching factory language.",
            "Use a top bar for factory context, week/date range, plan version, and sync status.",
            "Place global filters above work surfaces: buyer, order, style, workcenter, week, risk, owner.",
            "Open record details in right drawers so users retain table or timeline context.",
        ]
    )
    add_para(doc, "Status language", style="Heading 2")
    add_label_table(
        doc,
        [
            ("Status", "UI treatment and behavior"),
            ("Ready / Green", "Subtle green chip. Action is allowed, but details remain inspectable."),
            ("Conditional / Yellow", "Amber chip. Action requires owner, reason, due date, and risk note."),
            ("Blocked / Red", "Red chip. Release and freeze actions disabled until blocker is resolved or override is approved."),
            ("Black / Critical", "Dark chip. Management escalation and audit log required."),
            ("Stale data", "Muted badge with timestamp and source. Do not let stale data look safe."),
        ],
        widths=[Inches(1.55), Inches(4.85)],
        header=True,
    )
    add_para(doc, "Reusable front-end components", style="Heading 2")
    add_bullets(
        doc,
        [
            "Data grid with sticky header, frozen identity columns, row actions, column-level filters, and saved views.",
            "Timeline/Gantt lane with conflict overlays, frozen-zone indicators, and drag/edit impact preview.",
            "Readiness checklist with pass/fail/conditional states, owner assignment, and approval workflow.",
            "Capacity heatmap with keyboard-focusable cells, tooltips, and click-to-filter interactions.",
            "Exception drawer with owner, due date, suggested action, activity log, and closure evidence.",
            "Mobile task form with offline queue, local validation, sync status, and rejected-update retry path.",
        ]
    )
    doc.add_page_break()

    add_para(doc, "4. MVP Screen Blueprints", style="Heading 1")
    add_para(
        doc,
        "Each screen below includes a visual target and the front-end handoff contract needed to build the first implementation without losing the operating intent of the BRD.",
    )
    for idx, screen in enumerate(SCREENS, 1):
        add_para(doc, f"4.{idx} {screen.title}", style="Heading 2")
        add_mockup(doc, ASSET_DIR / screen.image, f"Figure {idx + 1} - {screen.title} visual composition.")
        add_screen_spec_table(doc, screen)
        if idx != len(SCREENS):
            doc.add_page_break()

    doc.add_page_break()
    add_para(doc, "5. Front-End Delivery Priorities", style="Heading 1")
    add_para(doc, "MVP implementation sequence", style="Heading 2")
    add_label_table(
        doc,
        [
            ("Phase", "Front-end scope"),
            ("1. Foundation shell", "Navigation, global filters, status chips, data grid, drawer framework, role-aware routing."),
            ("2. Planning core", "Order Lifecycle, PCD Readiness, Weekly Planning, Daily Release."),
            ("3. Capacity core", "Workcenter Load, Sewing Line Loading, Wash Planning."),
            ("4. Flow and recovery", "WIP Queues, Exception Center, recovery action workflow."),
            ("5. Dispatch protection", "Shipment Readiness, shipment blocker escalation, split shipment tracking."),
            ("6. Shopfloor capture", "Mobile/tablet update path with offline queue and sync feedback."),
        ],
        widths=[Inches(1.55), Inches(4.85)],
        header=True,
    )
    add_para(doc, "Acceptance criteria for the front end", style="Heading 2")
    add_bullets(
        doc,
        [
            "A planner can identify the highest shipment risk, root blocker, owner, and next action in under one minute.",
            "A blocked order cannot be released from PCD or daily production without an explicit override path.",
            "A weekly plan cannot be frozen while unresolved workcenter overload or material blockers are present.",
            "Every critical exception has one owner, one due date, one current status, and a visible action history.",
            "Tables, heatmaps, timelines, and drawers remain usable with large order volumes and narrow laptop screens.",
            "Status color is never the only signal; all states have text labels and accessible names.",
        ]
    )
    add_para(doc, "Responsive and accessibility requirements", style="Heading 2")
    add_bullets(
        doc,
        [
            "Desktop is primary for planning; tablet is primary for shopfloor update; mobile is limited to task capture and approvals.",
            "Keyboard navigation must work for grids, heatmap cells, filter controls, drawers, and action menus.",
            "Use WCAG-conscious contrast for status chips and never rely on red/green alone.",
            "Long tables need pagination or virtualization, but totals and selected filters must remain visible.",
            "All dates, quantities, statuses, and owner names should have deterministic formatting to avoid ambiguity.",
        ]
    )

    add_para(doc, "6. Engineering Notes", style="Heading 1")
    add_label_table(
        doc,
        [
            ("Topic", "Recommendation"),
            ("State management", "Use route-level state for filters and selected records so links can reproduce an operational view."),
            ("Data freshness", "Show source and last-sync time on every planning-critical screen."),
            ("Authorization", "Use role-based action permissions. Hidden actions are acceptable only if read-only users still understand why they cannot act."),
            ("Auditability", "Plan changes, conditional releases, override approvals, and shipment splits require immutable audit metadata."),
            ("Performance", "Use server-side filtering, pagination or virtualization, and summary endpoints for control tower metrics."),
            ("Design tokens", "Centralize risk colors, spacing, table density, chip styles, and typography in a shared UI package."),
        ],
        widths=[Inches(1.55), Inches(4.85)],
        header=True,
    )

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build_docx()
    print(path)
